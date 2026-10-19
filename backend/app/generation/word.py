"""Word 成品写入与完整性检查（规格 US 19 / 22）。

写入采用调研 A2 的**两段式**：模板里已定义样式就按样式写（保留用户模板版式）；
样式未定义（陷阱 2：潜在样式 / 无实体定义 → `KeyError`）就回退到「直接给段落设字号 /
加粗」，并在产物元数据标 `template_style_applied: false`——降级是**可见的**，不静默。

完整性检查（规格 US 22「结构损坏的成品不会被交到我手上」）：写出后立刻用 python-docx
**回读产物**，断言章节数与标题级别映射正确；不合格则丢弃文件并抛错，不落库、不交付。
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Pt

from app.generation.plan_document import Block, heading_count, parse_plan_blocks

# 方案块 → 模板样式名。标题级别由 `plan_document` 夹在 1~3（内置模板定义的上限）。
_HEADING_STYLE = "Heading {level}"
_BODY_STYLE = "Normal"
_BULLET_STYLE = "List Bullet"
_NUMBERED_STYLE = "List Number"
_TABLE_STYLE = "Table Grid"

# 回退（样式缺失）时直接给段落设置的格式。
_FALLBACK_HEADING_SIZES = {1: 20, 2: 16, 3: 14}
_FALLBACK_BODY_SIZE = 11


class DeliverableIntegrityError(Exception):
    """产物完整性检查未通过：结构损坏，不交付。"""


@dataclass(frozen=True)
class WrittenBlock:
    """一段已写入文档的内容及其所用样式（`None` = 走了回退的直接格式）。"""

    kind: str
    text: str
    style_name: str | None
    level: int = 0


@dataclass(frozen=True)
class WordBuild:
    """一次 Word 写入的结果与完整性检查所需的期望结构。"""

    path: Path
    blocks: tuple[Block, ...]
    written: tuple[WrittenBlock, ...]
    missing_styles: tuple[str, ...]
    base_paragraphs: int
    base_tables: int

    @property
    def template_style_applied(self) -> bool:
        """是否所有需要的样式都由模板定义（否则走了两段式回退）。"""
        return not self.missing_styles

    @property
    def headings(self) -> int:
        """产物应含的标题数。"""
        return heading_count(list(self.blocks))


def build_word_deliverable(
    plan_md: str, template_path: str | Path, output_path: str | Path
) -> WordBuild:
    """按模板把方案正文写成 .docx（原子落盘）。

    模板正文里已有的内容（封面 / 页眉页脚 / 节属性）随包保留——ADR-0003「不承诺无损」，
    但能留的样式与结构都留住。本函数只负责写；完整性检查由 `verify_word_deliverable` 做。
    """
    document = Document(str(template_path))
    blocks = parse_plan_blocks(plan_md)

    base_paragraphs = len(document.paragraphs)
    base_tables = len(document.tables)
    missing: set[str] = set()
    written: list[WrittenBlock] = []

    def style_available(name: str) -> bool:
        """模板是否已定义该样式（陷阱 2 的写前探测）。"""
        if name in document.styles:
            return True
        missing.add(name)
        return False

    for block in blocks:
        if block.kind == "heading":
            written.append(_write_heading(document, block, style_available))
        elif block.kind == "bullet":
            written.append(
                _write_styled_paragraph(
                    document, block.text, _BULLET_STYLE, fallback_prefix="• ", style_available=style_available
                )
            )
        elif block.kind == "numbered":
            written.append(
                _write_styled_paragraph(
                    document, block.text, _NUMBERED_STYLE, fallback_prefix="", style_available=style_available
                )
            )
        elif block.kind == "table":
            written.append(_write_table(document, block, style_available))
        else:
            written.append(
                _write_styled_paragraph(
                    document, block.text, _BODY_STYLE, fallback_prefix="", style_available=style_available
                )
            )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    # 先写临时文件再原子替换：中途失败不会留下半成品被当成成品文件。
    temp_path = output.with_name(output.name + ".part")
    document.save(str(temp_path))
    os.replace(temp_path, output)

    return WordBuild(
        path=output,
        blocks=tuple(blocks),
        written=tuple(written),
        missing_styles=tuple(sorted(missing)),
        base_paragraphs=base_paragraphs,
        base_tables=base_tables,
    )


def _write_heading(document, block: Block, style_available) -> WrittenBlock:
    """写标题：优先模板的 `Heading N`；未定义则回退为加粗 + 字号。"""
    style_name = _HEADING_STYLE.format(level=block.level)
    if style_available(style_name):
        document.add_paragraph(block.text, style=style_name)
        return WrittenBlock(kind="heading", text=block.text, style_name=style_name, level=block.level)

    paragraph = document.add_paragraph()
    run = paragraph.add_run(block.text)
    run.bold = True
    run.font.size = Pt(_FALLBACK_HEADING_SIZES.get(block.level, _FALLBACK_HEADING_SIZES[3]))
    return WrittenBlock(kind="heading", text=block.text, style_name=None, level=block.level)


def _write_styled_paragraph(
    document, text: str, style_name: str, *, fallback_prefix: str, style_available
) -> WrittenBlock:
    """写正文 / 列表项：优先模板样式；未定义则回退为无样式的直接段落。"""
    if style_available(style_name):
        document.add_paragraph(text, style=style_name)
        return WrittenBlock(kind="paragraph", text=text, style_name=style_name)

    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{fallback_prefix}{text}")
    run.font.size = Pt(_FALLBACK_BODY_SIZE)
    return WrittenBlock(kind="paragraph", text=f"{fallback_prefix}{text}", style_name=None)


def _write_table(document, block: Block, style_available) -> WrittenBlock:
    """写表格：优先模板的 `Table Grid`；未定义则回退为无边框表格。"""
    rows = block.rows
    widths = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=widths)

    style_name = _TABLE_STYLE
    if style_available(style_name):
        table.style = style_name
    else:
        style_name = None

    for row_index, row in enumerate(rows):
        for column_index, cell in enumerate(row):
            table.cell(row_index, column_index).text = cell

    flat = tuple(cell for row in rows for cell in row)
    return WrittenBlock(kind="table", text="|".join(flat), style_name=style_name)


def verify_word_deliverable(build: WordBuild) -> None:
    """**回读产物**做完整性检查：可打开、章节数正确、标题级别映射正确。

    检查只针对**本次追加的内容**（`base_paragraphs` / `base_tables` 之后的尾部），
    这样用户模板自带的封面文案不会被误判为结构损坏。
    """
    try:
        reopened = Document(str(build.path))
    except Exception as exc:  # 产物打不开 → 结构损坏，不交付
        raise DeliverableIntegrityError(f"产物无法被文档库打开：{exc}") from exc

    paragraphs = list(reopened.paragraphs)
    expected_paragraphs = [block for block in build.written if block.kind != "table"]

    if len(paragraphs) - build.base_paragraphs != len(expected_paragraphs):
        raise DeliverableIntegrityError(
            "产物段落数不符："
            f"期望 {len(expected_paragraphs)}，实际 {len(paragraphs) - build.base_paragraphs}"
        )

    appended = paragraphs[build.base_paragraphs:]
    for paragraph, expected in zip(appended, expected_paragraphs, strict=True):
        if expected.style_name is None:
            _verify_fallback(paragraph, expected)
        elif paragraph.style.name != expected.style_name:
            raise DeliverableIntegrityError(
                f"样式映射不符：期望 {expected.style_name}，实际 {paragraph.style.name}"
            )
        if paragraph.text != expected.text:
            raise DeliverableIntegrityError(
                f"内容不符：期望 {expected.text!r}，实际 {paragraph.text!r}"
            )

    tables = list(reopened.tables)
    expected_tables = [block for block in build.written if block.kind == "table"]
    if len(tables) - build.base_tables != len(expected_tables):
        raise DeliverableIntegrityError(
            f"产物表格数不符：期望 {len(expected_tables)}，实际 {len(tables) - build.base_tables}"
        )

    for table, expected in zip(tables[build.base_tables:], expected_tables, strict=True):
        flat = [cell.text for row in table.rows for cell in row.cells]
        if flat != expected.text.split("|"):
            raise DeliverableIntegrityError(f"表格内容不符：实际 {flat}")


def _verify_fallback(paragraph, expected: WrittenBlock) -> None:
    """回退分支的映射校验：标题必须是加粗 + 设过的字号。"""
    if expected.kind != "heading":
        return
    runs = paragraph.runs
    if not runs or not runs[0].bold:
        raise DeliverableIntegrityError(f"回退标题未加粗：{expected.text!r}")
    if runs[0].font.size != Pt(_FALLBACK_HEADING_SIZES.get(expected.level, 14)):
        raise DeliverableIntegrityError(f"回退标题字号不符：{expected.text!r}")