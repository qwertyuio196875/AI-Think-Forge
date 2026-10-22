"""把逐页 PNG 装配成满幅 PPTX 成品（工单 29 / ADR-0004）。

PPT 是**截图版**：先由 #28 截图渲染器把每一页 HTML 幻灯片截成 2560×1440 PNG，
本模块把它们按序铺进 13.333×7.5 英寸的 PPTX，一页一张满幅图（`add_picture` 铺满，
空版式，无文字占位）。尺寸换算用 python-pptx 的 `Inches`，像素→英寸的对应关系
（2560px/1440px @ 192dpi = 13.333in/7.5in）由 ADR-0004 固定。

两条硬纪律：
1. **一页一张、满幅、顺序由页序文件序决定**：每张 PNG 平铺整页，不缩放留白；
2. **任一页 PNG 尺寸不符（不是 2560×1440）或读不出时抛明确异常**，且在抛出前
   **不落盘**（先校验全部 PNG、内存里建完演示文稿最后才原子写文件），保证验收
   「磁盘上不出现 pptx 文件」。服务层编排里若中途失败还会主动删除半成品。
"""

from __future__ import annotations

import os
from dataclasses import dataclass
from pathlib import Path

from PIL import Image, UnidentifiedImageError
from pptx import Presentation
from pptx.util import Inches

# ADR-0004：截图固定 2560×1440 PNG，嵌入 pptx 时 13.333×7.5 英寸满幅。
IMAGE_WIDTH = 2560
IMAGE_HEIGHT = 1440

SLIDE_WIDTH_INCHES = 13.333
SLIDE_HEIGHT_INCHES = 7.5

# python-pptx 尺寸用 Inches；空版式（无文字占位，只放图片）。
_BLANK_LAYOUT = 6


class PptxAssemblyError(Exception):
    """PPTX 装配失败：任一页尺寸不符 / 读不出，不交付半成品。"""


@dataclass(frozen=True)
class PptxVerdict:
    """装配后的 PPTX 完整性检查结果。"""

    path: Path
    slide_count: int
    integrity_ok: bool = True


def _validated_pngs(slide_pngs) -> list[tuple[int, Path]]:
    """校验全部 PNG：必须是可读且尺寸严格等于 2560×1440。

    为什么先验全部再装配：任一页读不出 / 尺寸不符就整份作废，绝不落盘半成品。
    """
    validated: list[tuple[int, Path]] = []
    for index, png in slide_pngs:
        path = Path(png)
        try:
            with Image.open(path) as image:
                width, height = image.size
        except (UnidentifiedImageError, OSError) as exc:
            raise PptxAssemblyError(
                f"第 {index} 页截图不是可读图片：{path.name}（{exc}）"
            ) from exc
        if (width, height) != (IMAGE_WIDTH, IMAGE_HEIGHT):
            raise PptxAssemblyError(
                f"第 {index} 页截图尺寸不符：期望 {IMAGE_WIDTH}×{IMAGE_HEIGHT}，"
                f"实际 {width}×{height}"
            )
        validated.append((index, path))
    return validated


def assemble_pptx(
    slide_pngs: tuple[tuple[int, Path], ...], output_path: str | Path
) -> Path:
    """把 `(页序, PNG 路径)` 序列装配成满幅 PPTX 并落盘。

    返回落盘路径。任一页尺寸不符 / 读不出 → 抛 `PptxAssemblyError`，且**不落盘**。
    """
    validated = _validated_pngs(slide_pngs)

    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_WIDTH_INCHES)
    presentation.slide_height = Inches(SLIDE_HEIGHT_INCHES)

    for _index, png in validated:
        slide = presentation.slides.add_slide(presentation.slide_layouts[_BLANK_LAYOUT])
        slide.shapes.add_picture(
            str(png),
            Inches(0),
            Inches(0),
            width=presentation.slide_width,
            height=presentation.slide_height,
        )

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    # 先写临时文件再原子替换：中途失败不会留下半成品被当成成品文件。
    temp_path = output.with_name(output.name + ".part")
    presentation.save(str(temp_path))
    os.replace(temp_path, output)
    return output


def verify_pptx_deliverable(output_path: str | Path, *, expected_slides: int) -> PptxVerdict:
    """**回读产物**做完整性检查：能打开、页数等于规划页数、每页一张满幅图。

    任一条不通过抛 `PptxAssemblyError`（服务层据此丢弃产物、不落库、不交付）。
    pptx 满幅铺图不加文字框、不留白，故图片必须铺满整页。
    """
    path = Path(output_path)
    try:
        presentation = Presentation(str(path))
    except Exception as exc:  # 打不开 = 结构损坏
        raise PptxAssemblyError(f"成品 PPTX 无法被 python-pptx 打开：{exc}") from exc

    slide_count = len(presentation.slides)
    if slide_count != expected_slides:
        raise PptxAssemblyError(
            f"成品页数不符：期望 {expected_slides}，实际 {slide_count}"
        )

    for index, slide in enumerate(presentation.slides, start=1):
        pictures = [
            shape for shape in slide.shapes if shape.shape_type == 13  # MSO 图片
        ]
        if len(pictures) != 1:
            raise PptxAssemblyError(f"第 {index} 页图片数不符：期望 1，实际 {len(pictures)}")
        pic = pictures[0]
        if (
            abs(pic.left) != 0
            or abs(pic.top) != 0
            or abs(pic.width - presentation.slide_width) > 1000
            or abs(pic.height - presentation.slide_height) > 1000
        ):
            raise PptxAssemblyError(f"第 {index} 页不是满幅图片（未铺满整页）")

    return PptxVerdict(path=path, slide_count=slide_count)