"""文本与版式类资料的解析（TXT / MD / DOCX / PDF / PPTX）：正文、标题与基本结构。

选型与判据来自调研笔记 A1 / A5（`02-parsing-and-opendesign.md`）：

- TXT 编码顺序 `utf-8-sig` → `gb18030` → `charset-normalizer`，并把实际探测到的
  编码回传（前端可提示「检测到 GB18030」）。先严格解码、失败后才上探测器。
- MD 标题走 `markdown-it-py` 的 token 流，**不用正则**——正则会把代码围栏内的
  `#` 误判为标题。
- DOCX 用 `python-docx` 读 `paragraph.style.name` 判标题；作者手写加粗放大而不用
  内置样式时退回启发式，并在解析产物里标注 `heading_source`（style / heuristic / none）。
  图片、SmartArt、图表内文字、文本框浮动内容不进 `paragraphs`，读不到 → 归入未识别。
- PDF（TB-8）用 `pypdf`：逐页提取文本行，按**字号**启发式识别标题（PDF 没有样式信息，
  故不出现 `style` 来源）；页资源里的图片计入未识别。扫描件 / 图片型 PDF 没有文本层
  → 整份 `unrecognized`（`reason="no_text_layer"`、`needs_ocr=true`），**不做 OCR**。
- PPTX（TB-8）用 `python-pptx`：每页标题占位符（`heading_level=0`）、正文占位符与文本框
  逐段成块，表格**逐行转文本**（与 DOCX 同一口径）；图片 / 图表 / 组合形状读不到 → 逐项明示。

本模块只负责「能不能从文件里读出结构化的块」；「读不出来」一律显式返回
`unrecognized`，**绝不用空字符串冒充解析成功**（调研 A6）。
"""

from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

import charset_normalizer
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from markdown_it import MarkdownIt
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pypdf import PdfReader
from pypdf.errors import PdfReadError

# 本单接受的后缀（URL 属 TB-9）。
SUPPORTED_SUFFIXES = frozenset({".txt", ".md", ".docx", ".pdf", ".pptx"})

# 解析器标识：落 `documents.parser`，让「内容出自哪个解析器」可回溯。
PARSER_PLAIN_TEXT = "plain-text"
PARSER_DOCX = "python-docx"
PARSER_PDF = "pypdf"
PARSER_PPTX = "python-pptx"

# 编码探测顺序（A5）：先严格解码，失败后才让探测器兜底。
STRICT_ENCODINGS = ("utf-8-sig", "gb18030")

# 启发式标题判据（A1 边界 2）：短、加粗、显式字号且不以句读结尾。
_HEURISTIC_MAX_CHARS = 80
_HEURISTIC_MIN_PT = 14.0
_SENTENCE_ENDINGS = tuple("。！？；，,.;:!?")
_HEADING_STYLE = re.compile(r"heading\s+([1-9])", re.IGNORECASE)

# 版式类（PDF / PPTX）的标题启发式：字号明显大于正文字号、短、不以句读结尾。
_LAYOUT_HEADING_MIN_RATIO = 1.2
_LAYOUT_TITLE_RATIO = 1.8
_LAYOUT_SUBTITLE_RATIO = 1.4


class ParseError(ValueError):
    """文件无法解析（可读原因；调用方据此给出明确提示，不丢已存数据）。"""


@dataclass(frozen=True)
class ParsedBlock:
    """解析出的一个内容块（正文段 / 标题 / 表格行 / 代码块）。"""

    text: str
    # None = 正文；0 = Title（A1：level 0 是 Title）；1~9 = Heading N。
    heading_level: int | None = None
    # 定位信息（写进 `document_chunks.locator_json` 与总结引用）。
    locator: dict[str, Any] = field(default_factory=dict)
    heading_source: str | None = None


@dataclass(frozen=True)
class ParsedDocument:
    """一份资料的解析结果。`blocks` 为空时必须由调用方转成 `unrecognized`。"""

    format: str
    blocks: list[ParsedBlock]
    encoding: str | None = None
    heading_source: str = "none"
    images: int = 0
    tables: int = 0
    # 解析器标识（`pypdf` / `python-pptx` / `python-docx` / `plain-text`）。
    parser: str = PARSER_PLAIN_TEXT
    # 没有可分析文本时的原因；为 None 时由 `images` 推断（见 service）。
    empty_reason: str | None = None
    # 解析器逐项明示的「有正文但读不到」内容（含页 / 幻灯片级 locator）。
    unread_parts: list[dict[str, Any]] = field(default_factory=list)


def parse_document(path: str | Path, file_format: str) -> ParsedDocument:
    """按格式解析文件；读不出来抛 `ParseError`（不是「空内容成功」）。"""
    file_path = Path(path)
    suffix = f".{file_format.lower().lstrip('.')}"
    if suffix not in SUPPORTED_SUFFIXES:
        raise ParseError(f"本单不支持该格式：{suffix}")

    if suffix == ".docx":
        return _parse_docx(file_path)
    if suffix == ".pdf":
        return _parse_pdf(file_path)
    if suffix == ".pptx":
        return _parse_pptx(file_path)
    return _parse_text(file_path, suffix)


# --- TXT / MD ---------------------------------------------------------------


def _parse_text(path: Path, suffix: str) -> ParsedDocument:
    """TXT / MD：先解码（带编码回传），再按格式建块。"""
    raw = path.read_bytes()
    text, encoding = decode_text(raw)

    if suffix == ".md":
        blocks, headings = _parse_markdown(text)
        heading_source = "style" if headings else "none"
    else:
        blocks = _paragraph_blocks(text, "txt")
        heading_source = "none"

    return ParsedDocument(
        format=suffix.lstrip("."),
        blocks=blocks,
        encoding=encoding,
        heading_source=heading_source,
        parser=PARSER_PLAIN_TEXT,
    )


def decode_text(raw: bytes) -> tuple[str, str]:
    """按 A5 的顺序解码，返回 (文本, 实际编码名)。"""
    for encoding in STRICT_ENCODINGS:
        try:
            return raw.decode(encoding), encoding
        except UnicodeDecodeError:
            continue

    best = charset_normalizer.from_bytes(raw).best()
    if best is None or best.encoding is None:
        raise ParseError("无法识别文本编码；请把文件另存为 UTF-8 后重试")
    return str(best), best.encoding


def _paragraph_blocks(text: str, file_format: str) -> list[ParsedBlock]:
    """TXT：按行成块，跳过空行；位置用段落序号（1 起）。"""
    blocks: list[ParsedBlock] = []
    index = 0
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        index += 1
        blocks.append(
            ParsedBlock(
                text=stripped,
                locator={"format": file_format, "paragraph": index},
            )
        )
    return blocks


def parse_markdown_blocks(
    text: str,
    *,
    format_name: str = "md",
    locator_extra: dict[str, Any] | None = None,
) -> tuple[list[ParsedBlock], int]:
    """Markdown 文本 → 内容块 + 标题数（公开入口，供网页正文复用同一 token 流）。

    `format_name` / `locator_extra` 让调用方把格式名与来源地址（如最终 URL）写进
    每个块的 locator，使引用位置可回溯到具体来源。
    """
    return _parse_markdown(
        text, format_name=format_name, locator_extra=locator_extra
    )


def _parse_markdown(
    text: str,
    *,
    format_name: str = "md",
    locator_extra: dict[str, Any] | None = None,
) -> tuple[list[ParsedBlock], int]:
    """MD：走 token 流建块；代码围栏里的 `#` 不会被当成标题。

    标题级别取 `heading_open.tag`（h1..h6）；正文取 `inline`，代码取 `fence` /
    `code_block`（作为普通内容保留，不丢资料）。`token.map[0]` 是 0 起行号，
    没给时退回块序号，保证 locator 始终有可读位置。
    """
    tokens = MarkdownIt("commonmark").parse(text)
    blocks: list[ParsedBlock] = []
    headings = 0
    consumed: set[int] = set()
    block_index = 0

    for index, token in enumerate(tokens):
        if index in consumed:
            continue

        locator = _markdown_locator(
            token.map, block_index, format_name, locator_extra
        )
        if token.type == "heading_open":
            inline = tokens[index + 1] if index + 1 < len(tokens) else None
            if inline is not None and inline.type == "inline":
                consumed.add(index + 1)
            content = inline.content.strip() if inline is not None else ""
            if content:
                headings += 1
                block_index += 1
                blocks.append(
                    ParsedBlock(
                        text=content,
                        heading_level=int(token.tag.lstrip("h")),
                        locator=locator,
                        heading_source="style",
                    )
                )
        elif token.type in {"paragraph_open", "list_item_open"}:
            inline = tokens[index + 1] if index + 1 < len(tokens) else None
            if inline is None or inline.type != "inline":
                continue
            consumed.add(index + 1)
            content = inline.content.strip()
            if content:
                block_index += 1
                blocks.append(ParsedBlock(text=content, locator=locator))
        elif token.type in {"fence", "code_block"}:
            content = token.content.strip("\n").strip()
            if content:
                block_index += 1
                blocks.append(
                    ParsedBlock(
                        text=content,
                        locator={**locator, "code": True},
                    )
                )
        elif token.type == "inline":
            # 表格单元格等没有 paragraph_open 包裹的内联内容：不丢。
            content = token.content.strip()
            if content:
                block_index += 1
                blocks.append(ParsedBlock(text=content, locator=locator))

    return blocks, headings


def _markdown_locator(
    line_map: list[int] | None,
    fallback_index: int,
    format_name: str = "md",
    locator_extra: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """MD 定位：优先用 token.map 的起始行号（1 起），没有则退块序号。"""
    locator: dict[str, Any] = {"format": format_name, **(locator_extra or {})}
    if line_map:
        locator["line"] = line_map[0] + 1
    else:
        locator["block"] = fallback_index + 1
    return locator


# --- DOCX -------------------------------------------------------------------


def _parse_docx(path: Path) -> ParsedDocument:
    """DOCX：读段落 + 表格，图片数量单列（图片内文字读不到）。"""
    try:
        document = Document(str(path))
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as exc:
        raise ParseError(
            "文件已损坏或不是有效的 .docx（模板类型文件请另存为 .docx 后重试）"
        ) from exc
    except OSError as exc:
        raise ParseError(f"无法读取文件：{exc}") from exc

    blocks: list[ParsedBlock] = []
    heading_sources: set[str] = set()
    paragraph_index = 0

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        paragraph_index += 1
        level, source = _docx_heading(paragraph, text)
        if source is not None:
            heading_sources.add(source)
        blocks.append(
            ParsedBlock(
                text=text,
                heading_level=level,
                locator={"format": "docx", "paragraph": paragraph_index},
                heading_source=source,
            )
        )

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            row_text = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
            if not row_text:
                continue
            blocks.append(
                ParsedBlock(
                    text=row_text,
                    locator={
                        "format": "docx",
                        "table": table_index,
                        "row": row_index,
                    },
                )
            )

    if "style" in heading_sources:
        heading_source = "style"
    elif "heuristic" in heading_sources:
        heading_source = "heuristic"
    else:
        heading_source = "none"

    return ParsedDocument(
        format="docx",
        blocks=blocks,
        heading_source=heading_source,
        images=len(document.inline_shapes),
        tables=len(document.tables),
        parser=PARSER_DOCX,
    )


def _docx_heading(paragraph, text: str) -> tuple[int | None, str | None]:
    """判标题级别与来源：内置样式优先，其次字号 / 加粗启发式（A1）。"""
    style_name = paragraph.style.name if paragraph.style is not None else ""
    normalized = (style_name or "").strip()

    if normalized.lower() == "title":
        return 0, "style"
    matched = _HEADING_STYLE.fullmatch(normalized)
    if matched:
        return int(matched.group(1)), "style"

    return _heuristic_heading(paragraph, text)


def _heuristic_heading(paragraph, text: str) -> tuple[int | None, str | None]:
    """启发式：短、加粗、显式字号且不以句读结尾 → 标题（并标注来源）。"""
    if len(text) > _HEURISTIC_MAX_CHARS or text.endswith(_SENTENCE_ENDINGS):
        return None, None

    runs = [run for run in paragraph.runs if run.text.strip()]
    if not runs:
        return None, None

    first = runs[0]
    bold = first.bold
    size = first.font.size
    if bold is None:
        bold = paragraph.style.font.bold if paragraph.style is not None else None
    if size is None:
        size = paragraph.style.font.size if paragraph.style is not None else None

    if not bold or size is None or size.pt < _HEURISTIC_MIN_PT:
        return None, None

    if size.pt >= 18:
        level = 1
    elif size.pt >= 16:
        level = 2
    else:
        level = 3
    return level, "heuristic"


# --- PDF --------------------------------------------------------------------


def _parse_pdf(path: Path) -> ParsedDocument:
    """PDF：逐页提取文本行、按字号判标题；图片与「无文本层」如实标注。

    - 页 / 行级 locator：`{format: "pdf", page: N, line: M}`（均从 1 开始）；
    - 页面里的图片读不到 → `unread_parts` 逐页明示（`images_not_read`、`needs_ocr`）；
    - 整份没有文本层（扫描件 / 图片型）→ `empty_reason="no_text_layer"`，
      由 `service.primary_unrecognized` 短路，不调用模型。
    """
    try:
        reader = PdfReader(str(path))
        pages = [
            (page_number, page, _pdf_page_lines(page))
            for page_number, page in enumerate(reader.pages, start=1)
        ]
    except (PdfReadError, OSError, ValueError, TypeError, KeyError) as exc:
        raise ParseError(
            "文件已损坏、已加密或不是有效的 .pdf；请确认导出为 PDF 后重试"
        ) from exc

    blocks: list[ParsedBlock] = []
    unread_parts: list[dict[str, Any]] = []
    images = 0
    body_size = _pdf_body_font_size(pages)
    headings = 0

    for page_number, page, lines in pages:
        page_images = _pdf_page_images(page)
        if page_images:
            images += page_images
            unread_parts.append(
                {
                    "unrecognized": True,
                    "reason": "images_not_read",
                    "locator": {
                        "format": "pdf",
                        "page": page_number,
                        "images": page_images,
                    },
                    "needs_ocr": True,
                }
            )

        for line_number, (text, size) in enumerate(lines, start=1):
            level = _pdf_heading_level(size, body_size, text)
            if level is not None:
                headings += 1
            blocks.append(
                ParsedBlock(
                    text=text,
                    heading_level=level,
                    locator={
                        "format": "pdf",
                        "page": page_number,
                        "line": line_number,
                    },
                    heading_source="heuristic" if level is not None else None,
                )
            )

    return ParsedDocument(
        format="pdf",
        blocks=blocks,
        heading_source="heuristic" if headings else "none",
        images=images,
        parser=PARSER_PDF,
        empty_reason=None if blocks else "no_text_layer",
        unread_parts=unread_parts,
    )


def _pdf_page_lines(page) -> list[tuple[str, float | None]]:
    """页内文本行 + 该行最大字号。

    优先走 `extract_text(visitor_text=...)`：pypdf 逐段回调给出字号，据此重建「行」
    与行内最大字号（标题判定所需）；个别文本对象取不到字号时该行字号为 None。
    回调不可用（或该页提取失败）时退回整页纯文本，字号一律 None。
    """
    pieces: list[tuple[str, float | None]] = []

    def visitor(text, cm, tm, font_dict, font_size):  # pypdf 回调签名（未用到的形参保留）
        pieces.append((text, font_size))

    visited = True
    try:
        page.extract_text(visitor_text=visitor)
    except Exception:  # noqa: BLE001 - 单页提取失败不应让整份资料解析失败
        visited = False
        pieces = []

    lines = _lines_from_pieces(pieces)
    if lines or visited:
        return lines

    try:
        plain = page.extract_text() or ""
    except Exception:  # noqa: BLE001 - 同上
        return []
    return [(line.strip(), None) for line in plain.splitlines() if line.strip()]


def _lines_from_pieces(pieces: list[tuple[str, float | None]]) -> list[tuple[str, float | None]]:
    """把 pypdf 的逐段回调结果按换行折成行，并记录行内最大字号。"""
    lines: list[tuple[str, float | None]] = []
    buffer = ""
    sizes: list[float] = []

    for text, size in pieces:
        for index, part in enumerate(text.split("\n")):
            if index:
                lines.append((buffer.strip(), max(sizes) if sizes else None))
                buffer = ""
                sizes = []
            buffer += part
            if size is not None and part.strip():
                sizes.append(size)

    if buffer.strip():
        lines.append((buffer.strip(), max(sizes) if sizes else None))
    return [(text, size) for text, size in lines if text]


def _pdf_body_font_size(pages: list) -> float | None:
    """正文字号 = 全文出现最多的字号；只有一种字号时返回 None（无从区分标题）。"""
    sizes = [size for _, _, lines in pages for _, size in lines if size]
    if len(set(sizes)) < 2:
        return None
    return Counter(sizes).most_common(1)[0][0]


def _pdf_heading_level(size: float | None, body_size: float | None, text: str) -> int | None:
    """按字号比例判标题级别：短、不以句读结尾、明显大于正文。"""
    if size is None or body_size is None or not text:
        return None
    if len(text) > _HEURISTIC_MAX_CHARS or len(text) < 2:
        return None
    if text.endswith(_SENTENCE_ENDINGS):
        return None

    ratio = size / body_size
    if ratio < _LAYOUT_HEADING_MIN_RATIO:
        return None
    if ratio >= _LAYOUT_TITLE_RATIO:
        return 1
    if ratio >= _LAYOUT_SUBTITLE_RATIO:
        return 2
    return 3


def _pdf_page_images(page) -> int:
    """该页图片数量。

    pypdf 解码个别损坏图片会抛多种异常；此时退化为「页资源里的 XObject 数」——
    宁可保守明示（可能多报），也不把读不到的图静默丢掉。
    """
    try:
        return len(page.images)
    except Exception:  # noqa: BLE001 - 见上：解码失败退化为资源计数
        try:
            resources = page.get("/Resources") or {}
            return len(resources.get("/XObject") or {})
        except Exception:  # noqa: BLE001 - 资源也读不到就不计
            return 0


# --- PPTX -------------------------------------------------------------------


def _parse_pptx(path: Path) -> ParsedDocument:
    """PPTX：每页标题（title 占位符）/ 正文 / 表格逐行；图片、图表、组合形状明示未识别。

    - 幻灯片 / 形状级 locator：`{format: "pptx", slide: N, shape: M, ...}`（均从 1 开始），
      标题为 `{format: "pptx", slide: N, kind: "title"}`，表格行为
      `{format: "pptx", slide: N, table: T, row: R}`；
    - 组合形状（group）内的内容不递归展开：读不到就逐项明示，不假装已完整分析。
    """
    try:
        presentation = Presentation(str(path))
    except (PackageNotFoundError, BadZipFile, ValueError, KeyError) as exc:
        raise ParseError(
            "文件已损坏或不是有效的 .pptx（请另存为 .pptx 后重试）"
        ) from exc
    except OSError as exc:
        raise ParseError(f"无法读取文件：{exc}") from exc

    blocks: list[ParsedBlock] = []
    unread_parts: list[dict[str, Any]] = []
    images = 0
    tables = 0
    headings = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = _pptx_slide_title(slide)
        if title is not None:
            title_text = _pptx_shape_text(title)
            if title_text:
                headings += 1
                blocks.append(
                    ParsedBlock(
                        text=title_text,
                        heading_level=0,
                        locator={
                            "format": "pptx",
                            "slide": slide_number,
                            "kind": "title",
                        },
                        heading_source="style",
                    )
                )

        table_index = 0
        for shape_index, shape in enumerate(slide.shapes, start=1):
            if title is not None and shape.shape_id == title.shape_id:
                continue

            if getattr(shape, "has_table", False):
                table_index += 1
                tables += 1
                for row_index, row in enumerate(shape.table.rows, start=1):
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells
                    ).strip(" |")
                    if not row_text:
                        continue
                    blocks.append(
                        ParsedBlock(
                            text=row_text,
                            locator={
                                "format": "pptx",
                                "slide": slide_number,
                                "table": table_index,
                                "row": row_index,
                            },
                        )
                    )
                continue

            shape_locator = {
                "format": "pptx",
                "slide": slide_number,
                "shape": shape_index,
            }
            shape_type = getattr(shape, "shape_type", None)
            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                images += 1
                unread_parts.append(
                    _pptx_unread_part("images_not_read", shape_locator, needs_ocr=True)
                )
                continue
            if getattr(shape, "has_chart", False):
                unread_parts.append(
                    _pptx_unread_part("chart_not_read", shape_locator)
                )
                continue
            if shape_type == MSO_SHAPE_TYPE.GROUP:
                unread_parts.append(
                    _pptx_unread_part("grouped_shapes_not_read", shape_locator)
                )
                continue
            if getattr(shape, "has_text_frame", False):
                for paragraph_index, paragraph in enumerate(
                    shape.text_frame.paragraphs, start=1
                ):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    blocks.append(
                        ParsedBlock(
                            text=text,
                            locator={
                                **shape_locator,
                                "paragraph": paragraph_index,
                            },
                        )
                    )
                continue
            # 其他浮动对象（OLE / 嵌入对象等）读不到 → 明示。
            unread_parts.append(
                _pptx_unread_part("embedded_object_not_read", shape_locator)
            )

    return ParsedDocument(
        format="pptx",
        blocks=blocks,
        heading_source="style" if headings else "none",
        images=images,
        tables=tables,
        parser=PARSER_PPTX,
        empty_reason=None if blocks else "no_text_layer",
        unread_parts=unread_parts,
    )


def _pptx_slide_title(slide):
    """幻灯片标题占位符；没有标题占位符（空白版式）时返回 None。"""
    try:
        return slide.shapes.title
    except (ValueError, KeyError):
        return None


def _pptx_shape_text(shape) -> str:
    """形状里的可读文本（多段按行拼接）。"""
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(
        paragraph.text.strip()
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text.strip()
    )


def _pptx_unread_part(reason: str, locator: dict[str, Any], *, needs_ocr: bool = False) -> dict:
    """PPTX 的未识别条目（统一契约 + 幻灯片级 locator）。"""
    return {
        "unrecognized": True,
        "reason": reason,
        "locator": locator,
        "needs_ocr": needs_ocr,
    }
