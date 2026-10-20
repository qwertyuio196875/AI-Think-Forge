"""TB-5 可复现演示记录：生成真实 .docx 并回读章节与样式映射。

在仓库根运行（会真的写文件，产物落在 `data/deliverables/`，模板落在 `templates/`）：

    backend\\.venv\\Scripts\\python.exe backend\\scripts\\demo_word_deliverable.py

演示全程走**进程内 HTTP 接缝**（`TestClient`，不起端口、不联网），模型用内联脚本化
替身（零花费）——与自动化测试同一条路径，因此它既是演示也是证据链。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient

from app import config
from app.api.deps import get_model_client
from app.main import create_app

IDEA = "帮我准备一份面向新员工的产品培训"
PLAN_MD = (
    "# 结构化方案\n\n"
    "## 目标\n让新员工在入职两周内掌握产品全貌。\n\n"
    "## 读者\n新入职员工。\n\n"
    "## 章节\n### 产品定位\n从场景切入，而非功能罗列。\n\n"
    "### 核心场景\n- 早晨开工\n- 日常答疑\n\n"
    "## 关键约束\n篇幅控制在十页以内。\n\n"
    "| 项 | 值 |\n| --- | --- |\n| 篇幅 | 10 页 |\n"
)


class _ScriptedModelClient:
    """按顺序应答的模型客户端替身（演示不联网、零花费）。"""

    def __init__(self, *results: str) -> None:
        self._results = list(results)
        self._index = 0

    def complete(
        self,
        messages: list[dict],
        *,
        model: str | None = None,
        max_tokens: int | None = None,
        timeout: float | None = None,
    ) -> str:
        index = min(self._index, len(self._results) - 1)
        self._index += 1
        return self._results[index]


def _frontier_json(*triples: tuple[str, str, str]) -> str:
    return json.dumps(
        {
            "frontier": [
                {"title": title, "body": body, "recommended_answer": answer}
                for title, body, answer in triples
            ]
        },
        ensure_ascii=False,
    )


def _empty_frontier() -> str:
    return json.dumps({"frontier": []}, ensure_ascii=False)


def _print_mapping(label: str, path: str) -> None:
    """回读产物并打印章节与样式映射（标题级别 / 正文 / 列表 / 表格）。"""
    document = Document(path)
    print(f"\n[{label}] 回读 {path}")
    print(f"  标题数：{sum(1 for p in document.paragraphs if p.style.name.startswith('Heading '))}")
    for paragraph in document.paragraphs:
        if paragraph.style.name.startswith("Heading "):
            print(f"    {paragraph.style.name:<10} | {paragraph.text}")
        elif paragraph.style.name in ("Normal", "List Bullet"):
            print(f"    {paragraph.style.name:<10} | {paragraph.text[:28]}")
    for table in document.tables:
        rows = [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        print(f"    表格样式 {table.style.name}：{' / '.join(rows)}")


def _user_template(path: Path, *, heading1_pt: int) -> None:
    """造一份用户模板（换模板 → 样式随之变化）。"""
    document = Document()
    document.styles["Heading 1"].font.size = Pt(heading1_pt)
    path.write_bytes(b"")
    document.save(str(path))


def main() -> int:
    app = create_app(mount_frontend=False)
    # 同一个替身实例贯穿全部请求（依赖按请求解析，lambda 里新建实例会重置脚本）。
    model_client = _ScriptedModelClient(
        _frontier_json(("目标读者", "写给谁？", "新员工")), _empty_frontier(), PLAN_MD
    )
    app.dependency_overrides[get_model_client] = lambda: model_client

    templates_dir = Path(app.state.templates_dir)
    templates_dir.mkdir(parents=True, exist_ok=True)
    _user_template(templates_dir / "演示模板-大标题.docx", heading1_pt=28)

    with TestClient(app) as client:
        started = client.post("/api/grilling/tasks", json={"idea": IDEA}).json()
        task_id = started["task"]["id"]
        client.post(
            f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
        )
        client.post(f"/api/grilling/tasks/{task_id}/advance")
        plan = client.post(f"/api/grilling/tasks/{task_id}/plan/confirm").json()
        print(f"任务 {task_id}｜方案 {plan['id']}｜状态 {plan['status']}｜确认时间 {plan['confirmed_at']}")

        print("\n=== 1. 内置最小默认模板（开箱可用，未登记任何模板）===")
        builtin = client.post(f"/api/tasks/{task_id}/generations").json()
        print(json.dumps(builtin, ensure_ascii=False, indent=2))
        _print_mapping("内置默认模板", builtin["file_path"])

        print("\n=== 2. 任务级临时覆盖：换成用户模板，样式随模板变化 ===")
        print("登记：", client.post("/api/templates", json={"name": "演示模板-大标题.docx"}).json())
        overridden = client.post(
            f"/api/tasks/{task_id}/generations", json={"template": "演示模板-大标题.docx"}
        ).json()
        print("生成：", json.dumps(overridden, ensure_ascii=False, indent=2))
        _print_mapping("用户模板", overridden["file_path"])
        print(
            "产物的 Heading 1 字号："
            f"{Document(overridden['file_path']).styles['Heading 1'].font.size.pt} pt "
            f"（内置默认模板为 {Document(builtin['file_path']).styles['Heading 1'].font.size.pt} pt）"
        )
        listed = client.get("/api/templates").json()["items"]
        print(
            "模板清单（临时覆盖未改变全局默认）：",
            [(item["name"], item["is_default"], item["missing"]) for item in listed],
        )

        print("\n=== 3. 负测：模板缺失 / 非 .docx，明确报错且不产出成品 ===")
        missing = client.post(
            f"/api/tasks/{task_id}/generations", json={"template": "不存在的模板.docx"}
        )
        print("缺失模板 →", missing.status_code, missing.json()["detail"])
        (templates_dir / "演示模板.dotx").write_bytes(
            (templates_dir / "演示模板-大标题.docx").read_bytes()
        )
        dotx = client.post(f"/api/tasks/{task_id}/generations", json={"template": "演示模板.dotx"})
        print("dotx 模板 →", dotx.status_code, dotx.json()["detail"])

        print("\n=== 4. 陷阱 2：模板缺 Heading 样式 → 回退 + template_style_applied=false ===")
        stripped = templates_dir / "演示模板-无标题样式.docx"
        _strip_styles(templates_dir / "演示模板-大标题.docx", stripped)
        fallback = client.post(
            f"/api/tasks/{task_id}/generations", json={"template": "演示模板-无标题样式.docx"}
        ).json()
        print(
            "template_style_applied =",
            fallback["template_style_applied"],
            "｜missing_styles =",
            fallback["missing_styles"],
        )
        document = Document(fallback["file_path"])
        print(
            "回退后的标题（加粗直接格式）：",
            [p.text for p in document.paragraphs if p.runs and p.runs[0].bold][:3],
        )

        print(f"\n已生成成品记录：{client.get(f'/api/tasks/{task_id}/generations').json()['total']} 份")
        print(f"产物目录：{Path(config.DATA_DIR) / config.DELIVERABLES_SUBDIR / str(task_id)}")

    return 0


def _strip_styles(src: Path, dst: Path) -> None:
    """删掉一份 .docx 里的标题样式定义（模拟模板只留潜在样式，调研 A2 陷阱 2）。"""
    import zipfile
    from xml.etree import ElementTree

    from docx.styles import BabelFish

    internal = {
        BabelFish.ui2internal(name) for name in ("Heading 1", "Heading 2", "Heading 3")
    }
    with zipfile.ZipFile(src) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    root = ElementTree.fromstring(parts["word/styles.xml"])
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for style in list(root.findall(f"{w}style")):
        name_element = style.find(f"{w}name")
        if name_element is not None and name_element.get(f"{w}val") in internal:
            root.remove(style)
    parts["word/styles.xml"] = ElementTree.tostring(root, xml_declaration=True, encoding="UTF-8")

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


if __name__ == "__main__":
    raise SystemExit(main())