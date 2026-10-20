"""TB-7 可复现演示：长 DOCX → 分块 → 逐块分析 → 综合 → 引用位置 / 未识别。

在仓库根运行（会真的写文件，落在 `data/`）：

    backend\\.venv\\Scripts\\python.exe backend\\scripts\\demo_analysis_long_docx.py

演示全程走**进程内 HTTP 接缝**（`TestClient`，真实 multipart 上传 + 真实 SQLite /
文件落盘；不起端口、不联网），模型用内联脚本化替身（**零花费**）。因此它既是演示
也是证据链：分块数、逐块与综合调用次数、总结汉字数、引用位置、未识别提示都会打印。
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from docx import Document
from fastapi.testclient import TestClient

from app import config
from app.db.provider_repository import ProviderRepository
from app.main import create_app
from app.providers.model_client import ModelClientConfig
from app.providers.tiered_client import TieredModelClient

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
LONG_DOC_PARAGRAPHS = 320

# 一段「超长总结」：用来真实走一遍 500 汉字收敛（演示会打印 truncated 与最终字数）。
LONG_SUMMARY = (
    "本资料围绕产品定位、目标用户、核心场景与执行约束展开。"
    "作者主张以真实工作流而非功能清单来定义产品价值，并用可验证的小步交付降低风险。"
    "关键数据包括两周内完成首轮可用流程、单份资料总结不超过五百字、上传上限五十兆。"
    "结论是先把阅读到创作的链路跑通，再考虑跨文档检索。"
) * 6


class _NoKeyStore:
    """演示不触碰 Windows 凭据管理器；密钥缺失照常发请求的语义由生产实现负责。"""

    def get_key(self, provider_name: str) -> str | None:
        return None

    def set_key(self, provider_name: str, api_key: str) -> None:
        return None

    def delete_key(self, provider_name: str) -> None:
        return None

    def has_key(self, provider_name: str) -> bool:
        return False


class _ScriptedFactory:
    """固定返回同一个脚本化客户端，并记录解析到的 provider 配置。"""

    def __init__(self, client: _ScriptedClient) -> None:
        self.client = client
        self.configs: list[ModelClientConfig] = []

    def create(self, config: ModelClientConfig, *, api_key: str | None = None):
        self.configs.append(config)
        return self.client


class _ScriptedClient:
    """按提示词类型应答：逐块请求返回块分析，综合请求返回超长总结。"""

    def __init__(self) -> None:
        self.chunk_calls = 0
        self.reduce_calls = 0

    def complete(
        self,
        messages: list[dict],
        *,
        model: str | None = None,
        max_tokens: int | None = None,
        timeout: float | None = None,
    ) -> str:
        user_content = messages[-1]["content"]
        if "这是第" in user_content:
            self.chunk_calls += 1
            return (
                f"主题：第 {self.chunk_calls} 块的主题\n"
                "核心观点：围绕产品定位与用户场景展开\n"
                "关键内容：给出了可验证的执行约束与数据\n"
                "结论：本块结论明确"
            )
        self.reduce_calls += 1
        return LONG_SUMMARY


def _long_docx(path: Path) -> None:
    """造一份几百段的长 DOCX（标题 + 正文，足以切出多个分块）。"""
    document = Document()
    document.add_heading("长文资料：产品工作流设计", level=1)
    for index in range(1, LONG_DOC_PARAGRAPHS + 1):
        if index % 40 == 1:
            document.add_heading(f"第 {(index - 1) // 40 + 1} 章", level=2)
        document.add_paragraph(
            f"第 {index} 段：这一段描述产品定位、目标用户、核心场景与执行约束，"
            "强调先把阅读到创作的链路跑通，再用可验证的小步交付降低风险；"
            "同时记录重要数据与结论，确保长文在分块后不遗漏关键信息。"
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _image_only_docx(path: Path) -> None:
    """造一份只有图片、没有文本层的 DOCX（未识别 / 需 OCR 的演示样本）。"""
    import base64

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    document.add_picture(io.BytesIO(png))
    document.save(str(path))


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")

    config.DATA_DIR.mkdir(parents=True, exist_ok=True)
    demo_dir = config.DATA_DIR / "demo"
    db_path = demo_dir / "tb7-demo.db"
    demo_dir.mkdir(parents=True, exist_ok=True)
    db_path.unlink(missing_ok=True)

    long_path = demo_dir / "长文资料.docx"
    _long_docx(long_path)

    app = create_app(db_path=db_path, mount_frontend=False)
    scripted = _ScriptedClient()
    factory = _ScriptedFactory(scripted)
    app.state.tiered_model_client = TieredModelClient(
        ProviderRepository(db_path), _NoKeyStore(), factory
    )

    with TestClient(app) as client:
        print("=== 1. 长 DOCX 真实 multipart 上传（进程内 HTTP）===")
        upload = client.post(
            "/api/analysis/sources",
            files={"file": (long_path.name, long_path.read_bytes(), DOCX_MIME)},
        )
        uploaded = upload.json()
        document = uploaded["document"]
        chunk_count = document["chunks"]
        print(f"HTTP {upload.status_code}｜来源 {uploaded['source']['id']}｜任务 {uploaded['source']['task_id']}")
        print(
            f"parse_status={document['parse_status']}｜heading_source={document['heading_source']}"
            f"｜标题数={document['headings']}｜分块数={chunk_count}"
        )
        print(f"落盘位置：data/{uploaded['source']['ref']}")

        print("\n=== 2. 分析：逐块 + 综合（脚本化模型，零花费）===")
        analysis = client.post(f"/api/analysis/sources/{uploaded['source']['id']}/analyze")
        result = analysis.json()
        summary = result["summary"]
        print(f"HTTP {analysis.status_code}")
        print(
            f"逐块调用={scripted.chunk_calls}｜综合调用={scripted.reduce_calls}"
            f"｜总计={scripted.chunk_calls + scripted.reduce_calls}（应等于 分块数 {chunk_count} + 1）"
        )
        print(
            f"总结汉字数={summary['hanzi_count']}（上限 500）｜truncated={summary['truncated']}"
            f"｜chunk_count={summary['chunk_count']}"
        )
        print(
            f"档位：requested=turbo｜provider_kind={summary['provider_kind']}"
            f"｜provider_used={summary['provider_used']}｜model={summary['model_used']}"
            f"｜fallback={summary['provider_fallback']}"
        )
        print(f"总结前 80 字：{summary['text'][:80]}…")
        print(f"引用条数={len(summary['citations'])}")
        for citation in summary["citations"][:2] + summary["citations"][-1:]:
            print(
                f"  分块 {citation['chunk_seq']}｜位置 {citation['locator']}"
                f"｜摘录 {citation['excerpt'][:32]}…"
            )
        print(f"模型客户端收到的配置：{factory.configs[-1]}")

        print("\n=== 3. 一键送入主题创作（只预填，不调用模型）===")
        before = scripted.chunk_calls + scripted.reduce_calls
        sent = client.post(f"/api/analysis/sources/{uploaded['source']['id']}/send-to-creation").json()
        after = scripted.chunk_calls + scripted.reduce_calls
        dialogue = client.get(f"/api/grilling/tasks/{sent['task']['id']}/dialogue").json()
        print(
            f"新任务 {sent['task']['id']}｜kind={sent['task']['kind']}"
            f"｜已知信息含总结={summary['text'] in dialogue['seed_input']}"
            f"｜新增模型调用={after - before}"
        )

        print("\n=== 4. 未识别路径：只有图片、没有文本层 ===")
        image_path = demo_dir / "扫描件.docx"
        _image_only_docx(image_path)
        image_upload = client.post(
            "/api/analysis/sources",
            files={"file": (image_path.name, image_path.read_bytes(), DOCX_MIME)},
        ).json()
        print(
            f"parse_status={image_upload['document']['parse_status']}"
            f"｜chunks={image_upload['document']['chunks']}"
            f"｜unrecognized={image_upload['unrecognized']}"
        )
        calls_before = scripted.chunk_calls + scripted.reduce_calls
        image_analysis = client.post(
            f"/api/analysis/sources/{image_upload['source']['id']}/analyze"
        ).json()
        calls_after = scripted.chunk_calls + scripted.reduce_calls
        print(
            f"分析结果：summary={image_analysis['summary']}"
            f"｜unrecognized={image_analysis['unrecognized']}"
            f"｜新增模型调用={calls_after - calls_before}"
        )
        print(f"提示：{image_analysis['message']}")

        print("\n=== 5. 演示数据位置 ===")
        print(f"演示库：{db_path}")
        print(f"上传件目录：{config.DATA_DIR / 'uploads'}")
        print(f"长文样本：{long_path}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
