"""资料分析来源列表与历史续看入口（US16 / US26）接缝测试。"""

from __future__ import annotations

import io

from docx import Document
from fastapi.testclient import TestClient

from app.db.analysis_repository import AnalysisRepository

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _short_docx(text: str = "这是一份短资料，用于来源列表测试。") -> bytes:
    document = Document()
    document.add_heading("短资料", level=1)
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, name: str) -> dict:
    response = client.post(
        "/api/analysis/sources",
        files={"file": (name, _short_docx(), DOCX_MIME)},
    )
    assert response.status_code == 201
    return response.json()


def test_source_list_orders_by_id_desc_and_reports_summary_flag(
    analysis_client: TestClient,
) -> None:
    """来源列表按 id 倒序，带最新文档与总结标记。"""
    first = _upload(analysis_client, "第一份.docx")
    second = _upload(analysis_client, "第二份.docx")

    listed = analysis_client.get("/api/analysis/sources").json()

    assert listed["total"] == 2
    assert listed["limit"] == 20
    assert listed["offset"] == 0
    assert [item["source"]["id"] for item in listed["items"]] == [
        second["source"]["id"],
        first["source"]["id"],
    ]
    assert all(item["document"] is not None for item in listed["items"])
    assert all(item["has_summary"] is False for item in listed["items"])
    assert all(item["summary_id"] is None for item in listed["items"])
    assert all(item["summary_created_at"] is None for item in listed["items"])

    analyzed = analysis_client.post(
        f"/api/analysis/sources/{first['source']['id']}/analyze"
    )
    assert analyzed.status_code == 200
    summary = analyzed.json()["summary"]

    refreshed = analysis_client.get("/api/analysis/sources").json()["items"]
    by_id = {item["source"]["id"]: item for item in refreshed}
    assert by_id[first["source"]["id"]]["has_summary"] is True
    assert by_id[first["source"]["id"]]["summary_id"] == summary["id"]
    assert by_id[first["source"]["id"]]["summary_created_at"]
    assert by_id[second["source"]["id"]]["has_summary"] is False


def test_source_list_paginates_and_validates_parameters(
    analysis_client: TestClient,
) -> None:
    """分页 limit / offset 生效；limit 上限 100、offset 不可为负。"""
    uploaded = [_upload(analysis_client, f"资料{index}.docx") for index in range(3)]

    page = analysis_client.get("/api/analysis/sources", params={"limit": 2, "offset": 1}).json()

    assert page["total"] == 3
    assert page["limit"] == 2
    assert page["offset"] == 1
    assert [item["source"]["id"] for item in page["items"]] == [
        uploaded[1]["source"]["id"],
        uploaded[0]["source"]["id"],
    ]

    assert analysis_client.get("/api/analysis/sources", params={"limit": 101}).status_code == 422
    assert analysis_client.get("/api/analysis/sources", params={"offset": -1}).status_code == 422


def test_deleted_task_source_is_hidden_and_cannot_be_reanalyzed(
    analysis_client: TestClient,
) -> None:
    """软删除是留痕终态：来源不再出现在最近资料，也不能靠重新分析把任务改回 completed。"""
    uploaded = _upload(analysis_client, "已删除任务的资料.docx")
    source_id = uploaded["source"]["id"]
    task_id = uploaded["source"]["task_id"]

    deleted = analysis_client.delete(
        f"/api/tasks/{task_id}", params={"confirm": True}
    )
    assert deleted.status_code == 200

    listed = analysis_client.get("/api/analysis/sources").json()
    assert listed["total"] == 0
    assert all(item["source"]["id"] != source_id for item in listed["items"])

    analyzed = analysis_client.post(f"/api/analysis/sources/{source_id}/analyze")
    assert analyzed.status_code == 409
    assert "已删除" in analyzed.json()["detail"]

    # 状态没有因分析请求被改写
    assert analysis_client.get(f"/api/tasks/{task_id}").json()["status"] == "deleted"


def test_history_analysis_item_carries_latest_source_id(
    client: TestClient, db_path
) -> None:
    """分析历史条目带该任务最新的 source_id；创建类任务为 null。"""
    uploaded = _upload(client, "资料.docx")
    task_id = uploaded["source"]["task_id"]

    repository = AnalysisRepository(db_path)
    newer = repository.create_source(
        "file",
        "uploads/newer.txt",
        meta={"original_name": "更新的资料.txt"},
        task_id=task_id,
    )

    history = client.get("/api/history", params={"kind": "analysis"}).json()
    assert history["items"][0]["source_id"] == newer["id"]

    creation = client.post("/api/tasks", json={"kind": "creation", "title": "创作"}).json()
    creation_history = client.get("/api/history", params={"kind": "creation"}).json()
    assert creation_history["items"][0]["id"] == creation["id"]
    assert creation_history["items"][0]["source_id"] is None