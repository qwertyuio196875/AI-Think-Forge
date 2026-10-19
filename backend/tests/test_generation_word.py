"""TB-5 Word 成品（模板机制）接缝测试（唯一接缝 = 进程内 HTTP；模型走 fake；不联网）。

对应规格 User Stories 18~22 的 Word 一轨（ADR-0003 双轨之一）与工单 14 的窄路径：
- 正测：确认方案 → 生成 .docx → 产物文件存在且**可被文档库打开**；
  标题级别与正文样式映射正确（回读断言，不是只看文件存在）；
- 负测：模板缺失 / 损坏 → 明确报错且**不产生成品文件**（绝不静默退回空白生成）；
- `.dotx` / `.dotm` / `.docm` 与「内容类型不符的 .docx」（A2 陷阱 1）被拒；
- 模板样式未定义 → 回退两段式并标 `template_style_applied: false`（A2 陷阱 2）；
- 全局默认模板与任务级**临时覆盖**在数据与 API 上互不混淆；每次生成记录实际所用模板。

语义质量不做自动化断言（规格 Testing Decisions）——只断言状态、持久化与产物文件。
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from fakes import ScriptedModelClient
from fastapi.testclient import TestClient

IDEA = "帮我准备一份面向新员工的产品培训"

# 结构化方案草稿：含 H1/H2/H3、正文、列表、表格六类样式，用于验证样式映射。
PLAN_MD = (
    "# 结构化方案\n\n"
    "## 目标\n让新员工在入职两周内掌握产品全貌。\n\n"
    "## 读者\n新入职员工。\n\n"
    "## 章节\n### 产品定位\n从场景切入，而非功能罗列。\n\n"
    "### 核心场景\n- 早晨开工\n- 日常答疑\n\n"
    "## 关键约束\n篇幅控制在十页以内。\n\n"
    "| 项 | 值 |\n| --- | --- |\n| 篇幅 | 10 页 |\n"
)

EXPECTED_HEADINGS = {
    1: ["结构化方案"],
    2: ["目标", "读者", "章节", "关键约束"],
    3: ["产品定位", "核心场景"],
}


def _frontier_json(*triples: tuple[str, str, str]) -> str:
    """按输出契约拼装脚本化前沿应答。"""
    frontier = [
        {"title": title, "body": body, "recommended_answer": answer}
        for title, body, answer in triples
    ]
    return json.dumps({"frontier": frontier}, ensure_ascii=False)


def _confirmed_task(
    client: TestClient, scripted: ScriptedModelClient, plan_md: str = PLAN_MD
) -> int:
    """走完拷问 → 收敛 → 确认方案，返回已确认方案的任务 id。"""
    scripted.set_sequence(
        _frontier_json(("Q", "正文", "推荐")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        plan_md,
    )
    task_id = client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["答"]}
    )
    client.post(f"/api/grilling/tasks/{task_id}/advance")
    confirmed = client.post(f"/api/grilling/tasks/{task_id}/plan/confirm")
    assert confirmed.status_code == 200
    return task_id


def _headings(document) -> dict[int, list[str]]:
    """回读产物：按级别收集标题（`Heading N` 样式）。"""
    found: dict[int, list[str]] = {}
    for paragraph in document.paragraphs:
        name = paragraph.style.name
        if name.startswith("Heading "):
            found.setdefault(int(name.split()[-1]), []).append(paragraph.text)
    return found


# --- 正测：确认方案 → 生成可打开的 Word 成品 ------------------------------------


def test_confirmed_plan_produces_openable_word_deliverable(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """确认方案 → 生成 .docx；产物文件存在、可被 python-docx 打开、标题级别映射正确。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/generations")

    assert response.status_code == 201
    body = response.json()
    assert body["task_id"] == task_id
    assert body["format"] == "docx"
    assert body["integrity_ok"] is True
    assert body["template_style_applied"] is True

    artifact = Path(body["file_path"])
    assert artifact.exists()
    assert artifact.suffix == ".docx"

    document = Document(str(artifact))  # 可被文档库打开
    assert _headings(document) == EXPECTED_HEADINGS


def test_deliverable_maps_body_list_and_table_styles(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """正文 / 列表 / 表格三类样式同样按模板样式映射（不只看标题）。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)
    body = grilling_client.post(f"/api/tasks/{task_id}/generations").json()

    document = Document(body["file_path"])

    styles = {paragraph.text: paragraph.style.name for paragraph in document.paragraphs}
    assert styles["让新员工在入职两周内掌握产品全貌。"] == "Normal"
    assert styles["早晨开工"] == "List Bullet"
    assert styles["日常答疑"] == "List Bullet"

    assert len(document.tables) == 1
    table = document.tables[0]
    assert table.style.name == "Table Grid"
    assert [cell.text for cell in table.rows[0].cells] == ["项", "值"]
    assert [cell.text for cell in table.rows[1].cells] == ["篇幅", "10 页"]


def test_builtin_minimal_default_template_is_idempotent_and_complete(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
) -> None:
    """内置最小默认模板：首次运行幂等生成，至少定义六类样式。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)
    grilling_client.post(f"/api/tasks/{task_id}/generations")

    builtin = templates_dir / "_builtin" / "minimal-default.docx"
    assert builtin.exists()

    template = Document(str(builtin))
    for required in ("Heading 1", "Heading 2", "Heading 3", "Normal", "List Bullet", "Table Grid"):
        assert required in template.styles, f"内置模板缺少样式：{required}"

    stamp = builtin.stat().st_mtime_ns
    again = grilling_client.post(f"/api/tasks/{task_id}/generations")
    assert again.status_code == 201
    assert builtin.stat().st_mtime_ns == stamp  # 幂等：不重复生成


def test_generation_records_plan_format_path_and_actual_template(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """每次生成记录实际所用模板（规格 US 20）；记录可按任务读回。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)
    plan = grilling_client.get(f"/api/grilling/tasks/{task_id}/plan").json()

    created = grilling_client.post(f"/api/tasks/{task_id}/generations").json()

    assert created["plan_id"] == plan["id"]
    assert created["format"] == "docx"
    assert created["template_used"] == "_builtin/minimal-default.docx"
    assert created["created_at"]

    listed = grilling_client.get(f"/api/tasks/{task_id}/generations")
    assert listed.status_code == 200
    items = listed.json()["items"]
    assert [item["id"] for item in items] == [created["id"]]
    assert items[0]["template_used"] == "_builtin/minimal-default.docx"