"""任务级模型切换（US23）接缝测试。

唯一接缝 = 进程内 HTTP；主流程模型客户端使用**真实**的
`ProviderBackedModelClient` + 脚本化工厂，断言任务 pinned provider 的
`base_url` / `model_name` / 密钥确实抵达工厂。不联网、零真实花费。
"""

from __future__ import annotations

import io
import json

from docx import Document
from fakes import ScriptedModelClient, ScriptedModelClientFactory
from fastapi.testclient import TestClient

from app.api.deps import get_model_client
from app.db.message_repository import MessageRepository
from app.db.provider_repository import ProviderRepository
from app.grilling.service import SEED_ROUND_NO
from app.providers.provider_backed_client import ProviderBackedModelClient

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _provider(client: TestClient, kind: str) -> dict:
    items = client.get("/api/providers").json()["items"]
    return next(item for item in items if item["kind"] == kind)


def _short_docx() -> bytes:
    document = Document()
    document.add_heading("短资料", level=1)
    document.add_paragraph("这是一份用于模型覆盖测试的短资料。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _wire_real_client(app, key_store, factory) -> ProviderBackedModelClient:
    client = ProviderBackedModelClient(
        ProviderRepository(app.state.db_path), key_store, factory
    )
    app.dependency_overrides[get_model_client] = lambda: client
    return client


# --- 任务字段与 PUT /tasks/{id}/model -------------------------------------------


def test_create_task_with_provider_is_returned_and_persisted(client: TestClient) -> None:
    """创建任务可带 provider_id；响应与回读都带该字段。"""
    provider = _provider(client, "max")

    created = client.post(
        "/api/tasks",
        json={"kind": "creation", "title": "固定模型", "provider_id": provider["id"]},
    )
    assert created.status_code == 201
    assert created.json()["provider_id"] == provider["id"]

    read_back = client.get(f"/api/tasks/{created.json()['id']}").json()
    assert read_back["provider_id"] == provider["id"]

    plain = client.post("/api/tasks", json={"kind": "creation", "title": "默认模型"}).json()
    assert plain["provider_id"] is None


def test_put_task_model_sets_and_clears_override(client: TestClient) -> None:
    """PUT 设置 / 清除任务级模型覆盖；null 表示回到默认。"""
    provider = _provider(client, "max")
    task = client.post("/api/tasks", json={"kind": "creation", "title": "A"}).json()

    updated = client.put(f"/api/tasks/{task['id']}/model", json={"provider_id": provider["id"]})
    assert updated.status_code == 200
    assert updated.json()["provider_id"] == provider["id"]
    assert client.get(f"/api/tasks/{task['id']}").json()["provider_id"] == provider["id"]

    cleared = client.put(f"/api/tasks/{task['id']}/model", json={"provider_id": None})
    assert cleared.status_code == 200
    assert cleared.json()["provider_id"] is None


def test_put_task_model_rejects_missing_task_and_provider(client: TestClient) -> None:
    """任务不存在 404；provider 不存在 400（可读中文）。"""
    missing_task = client.put("/api/tasks/9999/model", json={"provider_id": None})
    assert missing_task.status_code == 404
    assert "任务 9999 不存在" in missing_task.json()["detail"]

    task = client.post("/api/tasks", json={"kind": "creation", "title": "A"}).json()
    missing_provider = client.put(
        f"/api/tasks/{task['id']}/model", json={"provider_id": 9999}
    )
    assert missing_provider.status_code == 400
    assert "模型配置 9999 不存在" in missing_provider.json()["detail"]


def test_put_task_model_rejects_deleted_task(client: TestClient) -> None:
    """已删除任务是留痕终态，不能再切模型（409）。"""
    task = client.post("/api/tasks", json={"kind": "creation", "title": "A"}).json()
    client.delete(f"/api/tasks/{task['id']}", params={"confirm": True})

    response = client.put(f"/api/tasks/{task['id']}/model", json={"provider_id": None})

    assert response.status_code == 409
    assert "已删除" in response.json()["detail"]


# --- 拷问主流程使用任务级 provider ------------------------------------------------


def test_grilling_start_uses_pinned_provider_config_and_key(
    client: TestClient, app, key_store
) -> None:
    """发起拷问时按指定 provider 的 base_url / model_name / 密钥调用模型。"""
    provider = _provider(client, "max")
    key_store.set_key(provider["name"], "sk-max")
    factory = ScriptedModelClientFactory(
        ScriptedModelClient(result=json.dumps({"frontier": []}, ensure_ascii=False))
    )
    _wire_real_client(app, key_store, factory)

    response = client.post(
        "/api/grilling/tasks",
        json={"idea": "写一份新员工手册", "provider_id": provider["id"]},
    )

    assert response.status_code == 201
    assert response.json()["task"]["provider_id"] == provider["id"]
    assert factory.last_config is not None
    assert factory.last_config.base_url == provider["base_url"]
    assert factory.last_config.model_name == provider["model_name"]
    assert factory.last_api_key == "sk-max"


def test_grilling_without_override_keeps_default_provider(
    client: TestClient, app, key_store
) -> None:
    """无覆盖时仍使用默认 provider（行为回归）。"""
    default = next(
        item
        for item in client.get("/api/providers").json()["items"]
        if item["is_default"]
    )
    factory = ScriptedModelClientFactory(
        ScriptedModelClient(result=json.dumps({"frontier": []}, ensure_ascii=False))
    )
    _wire_real_client(app, key_store, factory)

    response = client.post("/api/grilling/tasks", json={"idea": "普通任务"})

    assert response.status_code == 201
    assert response.json()["task"]["provider_id"] is None
    assert factory.last_config.model_name == default["model_name"]


def test_advance_and_revision_use_task_provider(
    client: TestClient, app, key_store, db_path
) -> None:
    """推进与方案修订都解析任务级 provider（不再一律用默认 provider）。"""
    provider = _provider(client, "max")
    factory = ScriptedModelClientFactory(
        ScriptedModelClient(result=json.dumps({"frontier": []}, ensure_ascii=False))
    )
    _wire_real_client(app, key_store, factory)

    task = client.post(
        "/api/tasks",
        json={"kind": "creation", "title": "固定模型", "provider_id": provider["id"]},
    ).json()
    # 补一条种子输入，使「无轮次推进」可据以提问（空前沿即收敛并生成方案）。
    MessageRepository(db_path).append(task["id"], SEED_ROUND_NO, "user", "写一份手册")
    advanced = client.post(f"/api/grilling/tasks/{task['id']}/advance")
    assert advanced.status_code == 200
    assert factory.last_config.model_name == provider["model_name"]

    revised = client.post(
        f"/api/grilling/tasks/{task['id']}/plan/revisions",
        json={"instruction": "补充执行步骤"},
    )
    assert revised.status_code == 200
    assert factory.last_config.model_name == provider["model_name"]
    assert factory.last_config.base_url == provider["base_url"]


# --- 资料分析使用任务级 provider ---------------------------------------------------


def test_analysis_uses_task_provider_instead_of_turbo(
    analysis_client: TestClient, scripted_model_client_factory: ScriptedModelClientFactory
) -> None:
    """来源所属任务 pinned provider 时，总结按该 provider 调用并如实记录。"""
    provider = _provider(analysis_client, "max")
    uploaded = analysis_client.post(
        "/api/analysis/sources",
        files={"file": ("短资料.docx", _short_docx(), DOCX_MIME)},
    ).json()
    task_id = uploaded["source"]["task_id"]

    pinned = analysis_client.put(
        f"/api/tasks/{task_id}/model", json={"provider_id": provider["id"]}
    )
    assert pinned.status_code == 200

    analyzed = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/analyze"
    )

    assert analyzed.status_code == 200
    summary = analyzed.json()["summary"]
    assert summary["provider_used"] == provider["name"]
    assert summary["provider_kind"] == provider["kind"]
    assert summary["model_used"] == provider["model_name"]
    assert summary["provider_fallback"] is False
    assert summary["provider_note"] == "任务级模型覆盖"
    assert scripted_model_client_factory.last_config.model_name == provider["model_name"]