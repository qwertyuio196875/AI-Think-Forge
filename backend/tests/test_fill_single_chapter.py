"""单章生成与真实正文 Word（工单 10 / ADR-0006；接缝 = 进程内 HTTP；模型走 fake）。

钉住工单 10 的四条验收：

1. **产出的 Word 里有真正的正文段落与小标题**，不是把方案要点原样搬过去；
2. **章节正文与方案里该章要点可对应**（要点进得了提示词，产出的是正文而不是提纲）；
3. **拼装后的文件能被 Word 打开**，结构完整（回读断言，不只看文件存在）；
4. **成品记录能区分「正文成品」与「大纲导出」**（`generations.source` 这一列现在预留好，
   文件名前缀的收口是工单 18 的事）。

本单**只生成第一章**：剩余章节属 #11、前情提要 #12、中断续写 #13、单章重写 #14、
字数分摊 #15、偏差修正 #16、长章分段 #17。逐章串行推进的编排留给 #11。

写作调用只发生在「写某一章」这个端点上；`POST /api/tasks/{id}/fill`（工单 9）只登记
填充任务、不调模型——那条既有契约不动，既有测试原样通过。
"""

from __future__ import annotations

import json
from pathlib import Path

import httpx
from docx import Document
from fakes import ScriptedModelClient
from fastapi.testclient import TestClient

from app.db.database import connect

IDEA = "帮我准备一份面向新员工的产品培训"

# plan-v3 形状的方案：两章，各带要点与建议字数。
PLAN_MD = (
    "# 结构化方案\n\n"
    "## 目标\n让新员工在入职两周内掌握产品全貌。\n\n"
    "## 读者\n新入职员工。\n\n"
    "## 章节\n"
    "- 标题：产品全景 ｜ 要点：从三个真实场景切入讲清主线 ｜ 建议字数：600\n"
    "- 标题：第一次上手 ｜ 要点：第一天要做的五件事 ｜ 建议字数：400\n\n"
    "## 关键约束\n篇幅控制在十页以内。\n"
)

# 注入给模型的「正文」：有小标题、有真正的叙述段落、有列表——与方案要点形态完全不同。
BODY_MD = (
    "## 早上开工的第一件事\n\n"
    "新员工到岗的第一个上午，其实不需要读完任何手册。"
    "真正要做的是把手从键盘上挪开一次，先看看别人每天是怎么用它的。\n\n"
    "- 打开工作台首页，看一眼今天的数据\n"
    "- 找一个同事问一句他昨天卡在哪里\n\n"
    "### 小结\n\n"
    "把这三件事做熟，产品在你眼里就不再是一堆功能，而是一段每天要走的流程。\n"
)

# 方案里第二章的标题：只写第一章时，它**不该**出现在正文成品里。
SECOND_CHAPTER_TITLE = "第一次上手"
FIRST_CHAPTER_TITLE = "产品全景"
FIRST_CHAPTER_POINTS = "从三个真实场景切入讲清主线"


def _frontier_json(*triples: tuple[str, str, str]) -> str:
    frontier = [
        {"title": title, "body": body, "recommended_answer": answer}
        for title, body, answer in triples
    ]
    return json.dumps({"frontier": frontier}, ensure_ascii=False)


def _task_with_started_fill(
    client: TestClient, scripted: ScriptedModelClient
) -> int:
    """走完拷问 → 收敛 → 确认方案 → 开始填充，返回任务 id。

    填充的写作应答是脚本里的最后一条（`set_sequence` 用尽后重复最后一条），
    因此本函数返回后，模型客户端的下一条应答就是 `BODY_MD`。
    """
    scripted.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
        BODY_MD,
    )
    task_id = client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    advanced = client.post(f"/api/grilling/tasks/{task_id}/advance")
    assert advanced.status_code == 200
    confirmed = client.post(f"/api/grilling/tasks/{task_id}/plan/confirm")
    assert confirmed.status_code == 200
    started = client.post(f"/api/tasks/{task_id}/fill", json={})
    assert started.status_code == 200
    return task_id


def _section_rows(db_path: Path, task_id: int) -> list[dict]:
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT seq, title, content_md, status, completed_at FROM plan_sections"
            " WHERE task_id = ? ORDER BY seq ASC",
            (task_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def _generation_rows(db_path: Path, task_id: int) -> list[dict]:
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT * FROM generations WHERE task_id = ? ORDER BY id ASC", (task_id,)
        ).fetchall()
    return [dict(row) for row in rows]


def _paragraphs(document, style_name: str) -> list[str]:
    """回读产物：按样式名收集段落文本。"""
    return [p.text for p in document.paragraphs if p.style.name == style_name]


def _all_text(document) -> str:
    return "\n".join(p.text for p in document.paragraphs)


# --- 验收 ①：正文落库 -----------------------------------------------------------


def test_filling_first_section_writes_body_and_marks_it_completed(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """写完第一章：正文落库、状态转「已完成」、完成时间有值；第二章仍待写。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 200
    body = response.json()
    assert body["task_id"] == task_id
    assert body["section"]["seq"] == 1
    assert body["section"]["title"] == FIRST_CHAPTER_TITLE
    assert body["section"]["status"] == "completed"
    assert body["section"]["status_label"] == "已完成"
    assert body["section"]["hanzi_count"] > 0
    assert body["section"]["completed_at"] is not None

    rows = _section_rows(db_path, task_id)
    assert rows[0]["content_md"] == BODY_MD.strip()
    assert rows[0]["status"] == "completed"
    assert rows[0]["completed_at"] is not None
    # 只写第一章：第二章一个字都没有。
    assert rows[1]["content_md"] == ""
    assert rows[1]["status"] == "pending"


def test_filled_section_is_visible_through_the_read_only_section_list(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """只读章节清单与填充响应是同一份事实（前端不必再解析一套）。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    filled = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill").json()
    listed = grilling_client.get(f"/api/tasks/{task_id}/sections").json()

    assert listed["items"][0] == filled["section"]
    assert listed["items"][1]["status"] == "pending"


def test_fill_progress_counts_completed_sections(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """填充任务的「已完成章节数」随写随更新（进度不只活在内存里）。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    assert grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill").status_code == 200

    with connect(db_path) as connection:
        row = connection.execute(
            "SELECT completed_sections FROM fill_tasks WHERE task_id = ? ORDER BY id DESC LIMIT 1",
            (task_id,),
        ).fetchone()
    assert int(row["completed_sections"]) == 1


# --- 验收 ②：正文与方案要点可对应（要点进得了提示词，产出的是正文）----------------


def test_writing_prompt_carries_the_chapter_title_and_key_points(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """本章标题与要点必须进提示词：正文才可能与方案里该章要点对应。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    assert grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill").status_code == 200

    sent = "\n".join(
        message["content"] for message in scripted_model_client.calls[-1]["messages"]
    )
    assert FIRST_CHAPTER_TITLE in sent
    assert FIRST_CHAPTER_POINTS in sent
    # 目标字数也进上下文（按汉字计）。
    assert "600" in sent


def test_writing_call_uses_per_call_token_cap_and_timeout(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """长章节用 #3 的按次覆盖：输出上限与超时都放宽，不必改客户端默认值。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    assert grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill").status_code == 200

    call = scripted_model_client.calls[-1]
    assert call["max_tokens"] == 8192
    assert call["timeout"] == 180.0


# --- 验收 ③：Word 里有真实正文段落与小标题，文件可被打开 -------------------------


def test_word_deliverable_contains_real_paragraphs_and_subheadings(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """产出的 .docx：章标题是小标题、正文是真段落、列表是列表项。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 200
    artifact = Path(response.json()["deliverable"]["file_path"])
    assert artifact.exists()
    assert artifact.suffix == ".docx"

    # 可被文档库打开，结构完整。
    document = Document(str(artifact))
    # 章标题是成品目录的一级标题（带章序号）；章内小标题降一级，成为二级 / 三级标题。
    assert "第 1 章 · 产品全景" in _paragraphs(document, "Heading 1")
    assert "早上开工的第一件事" in _paragraphs(document, "Heading 2")
    assert "小结" in _paragraphs(document, "Heading 3")
    assert any(
        "真正要做的是把手从键盘上挪开一次" in text
        for text in _paragraphs(document, "Normal")
    )
    assert "打开工作台首页，看一眼今天的数据" in _paragraphs(document, "List Bullet")

    # 不是把方案要点原样搬过去：要点那句话不在成品里。
    assert FIRST_CHAPTER_POINTS not in _all_text(document)
    # 只写了第一章，第二章不出现。
    assert SECOND_CHAPTER_TITLE not in _all_text(document)


def test_word_deliverable_file_name_has_no_outline_prefix(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """正文成品的文件名不带 `outline-` 前缀（大纲导出的收口是 #18 的事，但这里不能带）。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    name = Path(response.json()["deliverable"]["file_path"]).name
    assert not name.startswith("outline-")


# --- 验收 ④：成品记录能区分「正文成品」与「大纲导出」------------------------------


def test_generation_record_distinguishes_body_deliverable_from_outline_export(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """同一任务下的两份成品：`source` 一个是 body、一个是 outline，前端据此分辨。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    outline = grilling_client.post(f"/api/tasks/{task_id}/generations")
    assert outline.status_code == 201
    assert outline.json()["source"] == "outline"

    body = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")
    assert body.status_code == 200
    assert body.json()["deliverable"]["source"] == "body"

    listed = grilling_client.get(f"/api/tasks/{task_id}/generations").json()
    assert [item["source"] for item in listed["items"]] == ["outline", "body"]
    assert listed["items"][0]["source_label"] == "大纲导出"
    assert listed["items"][1]["source_label"] == "正文成品"

    rows = _generation_rows(db_path, task_id)
    assert [row["source"] for row in rows] == ["outline", "body"]


# --- 闸门与负测 ------------------------------------------------------------------


def test_writing_is_rejected_while_plan_is_not_confirmed(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """方案未确认 → 409，绝不写正文（填充是唯一会调模型的写作环节）。"""
    scripted_model_client.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
    )
    task_id = grilling_client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client = grilling_client
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    client.post(f"/api/grilling/tasks/{task_id}/advance")

    response = client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 409
    assert "确认" in response.json()["detail"]


def test_writing_unknown_section_returns_404(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """方案里没有这一章 → 404（不静默建行）。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/9/fill")

    assert response.status_code == 404
    assert "9" in response.json()["detail"]


def test_writing_requires_a_started_fill(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """没有进行中的填充任务 → 409：填充是独立阶段，得先「确认开始填充」。"""
    scripted_model_client.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
    )
    task_id = grilling_client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client = grilling_client
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    client.post(f"/api/grilling/tasks/{task_id}/advance")
    client.post(f"/api/grilling/tasks/{task_id}/plan/confirm")

    response = client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 409
    assert "填充" in response.json()["detail"]


def test_finished_section_cannot_be_written_again(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """已完成的章不重复写：重写是工单 14 的职责（另开端点）。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)
    assert grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill").status_code == 200

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 409
    assert "已完成" in response.json()["detail"]


def test_model_failure_marks_section_failed_and_produces_no_deliverable(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    db_path: Path,
    data_dir: Path,
) -> None:
    """模型调用失败 → 502、章节标「失败」、不产出任何成品文件。"""
    task_id = _task_with_started_fill(grilling_client, scripted_model_client)
    scripted_model_client.set_error(httpx.ConnectError("boom"))

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    assert response.status_code == 502
    assert "第 1 章" in response.json()["detail"]

    rows = _section_rows(db_path, task_id)
    assert rows[0]["status"] == "failed"
    assert rows[0]["content_md"] == ""
    assert _generation_rows(db_path, task_id) == []
    deliverables = data_dir / "deliverables" / str(task_id)
    assert not deliverables.exists() or list(deliverables.glob("*.docx")) == []
