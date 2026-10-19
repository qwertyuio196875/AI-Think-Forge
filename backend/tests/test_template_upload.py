"""TB-7 模板 multipart 上传（补齐约定笔记里归 TB-7 的「模板文件上传」）。

只接受 `.docx`，复用同一套扩展名 + 大小校验；`.dotx` 与改名成 `.docx` 的模板
类型文件沿用 TB-5 的错误口径（「请另存为 .docx」）。
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("模板标题", level=1)
    document.add_paragraph("模板正文。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_template_upload_accepts_docx_and_registers_it(
    client: TestClient, templates_dir: Path
) -> None:
    """合法 .docx 上传 → 落进 `templates/` 并登记，设置页无需手动放文件。"""
    response = client.post(
        "/api/templates/upload",
        files={"file": ("品牌模板.docx", _docx_bytes(), DOCX_MIME)},
    )

    assert response.status_code == 201
    body = response.json()
    assert body["template"]["name"] == "品牌模板.docx"
    assert body["template"]["registered"] is True
    assert body["template"]["usable"] is True
    assert (templates_dir / "品牌模板.docx").is_file()

    listed = client.get("/api/templates").json()["items"]
    entry = next(item for item in listed if item["name"] == "品牌模板.docx")
    assert entry["registered"] is True


def test_template_upload_rejects_non_docx(client: TestClient, templates_dir: Path) -> None:
    """`.dotx` → 400，并提示「请另存为 .docx」；文件不落盘。"""
    response = client.post(
        "/api/templates/upload",
        files={"file": ("品牌模板.dotx", _docx_bytes(), DOCX_MIME)},
    )

    assert response.status_code == 400
    assert "另存为 .docx" in response.json()["detail"]
    assert not (templates_dir / "品牌模板.dotx").exists()


def test_template_upload_rejects_corrupt_docx(
    client: TestClient, templates_dir: Path
) -> None:
    """改名为 `.docx` 的模板类型 / 损坏文件 → 400，且不留下半成品。"""
    response = client.post(
        "/api/templates/upload",
        files={"file": ("损坏.docx", b"not a zip package", DOCX_MIME)},
    )

    assert response.status_code == 400
    assert "另存为 .docx" in response.json()["detail"]
    assert not (templates_dir / "损坏.docx").exists()


def test_template_upload_does_not_silently_overwrite(
    client: TestClient, templates_dir: Path
) -> None:
    """同名模板已存在 → 409，不静默覆盖用户文件。"""
    first = client.post(
        "/api/templates/upload",
        files={"file": ("品牌模板.docx", _docx_bytes(), DOCX_MIME)},
    )
    assert first.status_code == 201

    second = client.post(
        "/api/templates/upload",
        files={"file": ("品牌模板.docx", _docx_bytes(), DOCX_MIME)},
    )

    assert second.status_code == 409
    assert "已存在" in second.json()["detail"]
