"""单章重写（工单 14 / ADR-0006；接缝 = 进程内 HTTP；模型走 fake）。

钉住工单 14 的四条验收：

1. **只有目标章正文改变，其他章正文逐字不变**；
2. **成品重新拼装，不产生第二个文件让用户自己合并**；
3. **成品总字数随重写结果变化**；
4. **可重写的目标限于已完成或失败的章**。
"""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from fakes import ScriptedModelClient
from fastapi.testclient import TestClient

from app.db.database import connect
from app.hanzi import count_hanzi

IDEA = "帮我准备一份面向新员工的产品培训"

PLAN_MD = (
    "# 结构化方案\n\n"
    "## 目标\n让新员工在入职两周内掌握产品全貌。\n\n"
    "## 章节\n"
    "- 标题：产品全景 ｜ 要点：从三个真实场景切入讲清主线 ｜ 建议字数：600\n"
    "- 标题：第一次上手 ｜ 要点：第一天要做的五件事 ｜ 建议字数：400\n"
)

BODY_FIRST = "## 情景\n\n第一篇正文，讲产品主线怎么串起来。\n"
BODY_SECOND = "## 五件事\n\n第二篇正文，讲第一天要做的五件事。\n"
# 重写后的第一章：明显更长，用来验证成品总字数随之变化。
BODY_FIRST_REWRITTEN = (
    "## 情景\n\n"
    "重写后的第一篇正文，讲产品主线从哪里来、往哪里去，"
    "以及新员工第一天为什么会觉得它像一堆功能。\n\n"
    "## 为什么这样安排\n\n"
    "因为先看流程再看功能，才不会把系统当成一个零件箱。\n"
)

_SUMMARY = "本章讲完了主线，给下一章留下了承接点。"


def _frontier_json(*triples: tuple[str, str, str]) -> str:
    frontier = [
        {"title": title, "body": body, "recommended_answer": answer}
        for title, body, answer in triples
    ]
    return json.dumps({"frontier": frontier}, ensure_ascii=False)


def _filled_task(
    client: TestClient, scripted: ScriptedModelClient, *bodies: str
) -> int:
    """走完拷问 → 确认 → 填充 → 跑完全篇，返回任务 id。"""
    responses: list[str] = []
    for body in bodies or (BODY_FIRST, BODY_SECOND):
        responses.extend([body, _SUMMARY])
    scripted.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
        *responses,
    )
    task_id = client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    assert client.post(f"/api/grilling/tasks/{task_id}/advance").status_code == 200
    assert client.post(f"/api/grilling/tasks/{task_id}/plan/confirm").status_code == 200
    assert client.post(f"/api/tasks/{task_id}/fill", json={}).status_code == 200
    assert client.post(f"/api/tasks/{task_id}/fill/run").status_code == 200
    return task_id


def _section_rows(db_path: Path, task_id: int) -> list[dict]:
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT seq, title, content_md, status FROM plan_sections"
            " WHERE task_id = ? ORDER BY seq ASC",
            (task_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def _body_generations(db_path: Path, task_id: int) -> list[dict]:
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT * FROM generations WHERE task_id = ? AND source = 'body' ORDER BY id ASC",
            (task_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def _all_generations(db_path: Path, task_id: int) -> list[dict]:
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT * FROM generations WHERE task_id = ? ORDER BY id ASC", (task_id,)
        ).fetchall()
    return [dict(row) for row in rows]


# --- 验收 ①：只有目标章改变 --------------------------------------------------------


def test_rewrite_changes_only_the_target_section(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """重写第一章：第二章正文逐字不变。"""
    task_id = _filled_task(grilling_client, scripted_model_client)
    before = _section_rows(db_path, task_id)
    scripted_model_client.set_sequence(BODY_FIRST_REWRITTEN)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    assert response.status_code == 200
    assert response.json()["section"]["seq"] == 1
    assert response.json()["section"]["status"] == "completed"

    after = _section_rows(db_path, task_id)
    assert after[0]["content_md"] == BODY_FIRST_REWRITTEN.strip()
    # 第二章一个字都不动。
    assert after[1]["content_md"] == before[1]["content_md"] == BODY_SECOND.strip()
    assert after[1]["status"] == "completed"


# --- 验收 ②：成品重新拼装，不产生第二个文件 ---------------------------------------


def test_rewrite_rebuilds_the_same_deliverable_instead_of_adding_a_second(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """重写后任务下**只有一份**正文成品记录，且文件是新的那份。"""
    task_id = _filled_task(grilling_client, scripted_model_client)
    before = _body_generations(db_path, task_id)
    assert len(before) == 1
    old_path = Path(before[0]["file_path"])
    assert old_path.exists()

    scripted_model_client.set_sequence(BODY_FIRST_REWRITTEN)
    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    assert response.status_code == 200
    after = _body_generations(db_path, task_id)
    assert len(after) == 1, "重写不该留下第二份正文成品"
    # 旧文件被清理，新文件存在。
    assert not old_path.exists()
    new_path = Path(after[0]["file_path"])
    assert new_path.exists()
    # 响应里的成品就是库里的那一份。
    assert response.json()["deliverable"]["id"] == after[0]["id"]


def test_rewrite_keeps_outline_export_untouched(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """大纲导出是另一条产物：重写正文不该把它删掉。"""
    task_id = _filled_task(grilling_client, scripted_model_client)
    assert grilling_client.post(f"/api/tasks/{task_id}/generations").status_code == 201
    assert len(_all_generations(db_path, task_id)) == 2  # 1 份大纲 + 1 份正文

    scripted_model_client.set_sequence(BODY_FIRST_REWRITTEN)
    grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    rows = _all_generations(db_path, task_id)
    assert [row["source"] for row in rows] == ["outline", "body"]


# --- 验收 ③：总字数随重写变化 ------------------------------------------------------


def test_total_hanzi_changes_with_the_rewrite(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """重写的正文更长 → 章节汉字数与成品总字数都随之变大。"""
    task_id = _filled_task(grilling_client, scripted_model_client)
    before_rows = _section_rows(db_path, task_id)
    before_total = sum(count_hanzi(row["content_md"]) for row in before_rows)

    scripted_model_client.set_sequence(BODY_FIRST_REWRITTEN)
    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    after_rows = _section_rows(db_path, task_id)
    after_total = sum(count_hanzi(row["content_md"]) for row in after_rows)
    assert after_total > before_total
    assert response.json()["section"]["hanzi_count"] == count_hanzi(
        BODY_FIRST_REWRITTEN
    )


# --- 验收 ④：只有已完成或失败的章可重写 -------------------------------------------


def test_pending_section_cannot_be_rewritten(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """还没写过的章 → 409（那是正常填充的活儿，不是重写）。"""
    # 只填第一章就停：第二章仍是 pending。
    scripted_model_client.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
        BODY_FIRST,
        _SUMMARY,
    )
    task_id = grilling_client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"][
        "id"
    ]
    grilling_client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    grilling_client.post(f"/api/grilling/tasks/{task_id}/advance")
    grilling_client.post(f"/api/grilling/tasks/{task_id}/plan/confirm")
    grilling_client.post(f"/api/tasks/{task_id}/fill", json={})
    grilling_client.post(f"/api/tasks/{task_id}/sections/1/fill")

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/2/rewrite")

    assert response.status_code == 409
    assert "还没有写过" in response.json()["detail"]


def test_unknown_section_returns_404(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """方案里没有这一章 → 404。"""
    task_id = _filled_task(grilling_client, scripted_model_client)

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/9/rewrite")

    assert response.status_code == 404
    assert "9" in response.json()["detail"]


def test_rewrite_is_rejected_while_plan_is_not_confirmed(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """方案未确认 → 409，绝不写正文。"""
    scripted_model_client.set_sequence(
        _frontier_json(("目标读者", "写给谁？", "新员工")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
    )
    task_id = grilling_client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"][
        "id"
    ]
    grilling_client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["新员工"]}
    )
    grilling_client.post(f"/api/grilling/tasks/{task_id}/advance")

    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    assert response.status_code == 409
    assert "确认" in response.json()["detail"]


def test_rewrite_failure_keeps_the_old_body(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """重写时模型失败 → 502，**旧正文与旧成品都还在**（不把已写好的章弄丢）。"""
    import httpx

    task_id = _filled_task(grilling_client, scripted_model_client)
    before = _section_rows(db_path, task_id)
    before_gens = _body_generations(db_path, task_id)

    scripted_model_client.set_error(httpx.ConnectError("rewrite boom"))
    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    assert response.status_code == 502
    assert "第 1 章" in response.json()["detail"]
    after = _section_rows(db_path, task_id)
    assert after[0]["content_md"] == before[0]["content_md"]
    assert after[0]["status"] == "completed"
    # 成品记录与文件都没动。
    assert _body_generations(db_path, task_id) == before_gens
    assert Path(before_gens[0]["file_path"]).exists()


def test_rewritten_deliverable_contains_the_new_text(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, db_path: Path
) -> None:
    """重新拼装的成品里能读到重写后的文字，且章节目录不变。"""
    task_id = _filled_task(grilling_client, scripted_model_client)

    scripted_model_client.set_sequence(BODY_FIRST_REWRITTEN)
    response = grilling_client.post(f"/api/tasks/{task_id}/sections/1/rewrite")

    artifact = Path(response.json()["deliverable"]["file_path"])
    document = Document(str(artifact))
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert headings == ["第 1 章 · 产品全景", "第 2 章 · 第一次上手"]
    text = "\n".join(p.text for p in document.paragraphs)
    assert "往哪里去" in text
