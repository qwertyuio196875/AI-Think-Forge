"""嵌入写入与向量可见状态（工单 20）接缝测试。

接缝只到进程内 HTTP + 数据库可见状态：上传一份资料，看 `embeddings` 表里是否
真的写了、写成了什么格式、失败时资料是否还读得到。嵌入客户端一律用
`FakeEmbeddingClient` 注入（确定性假向量、**绝不联网**），失败与维度不一致
同样靠替身构造。

不测实现细节：不断言「嵌入被调用了 N 次」，只断言库里能看见的结果。
"""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from fakes import FakeEmbeddingClient
from fastapi.testclient import TestClient

from app.api.deps import EmbeddingClientDep, get_embedding_client
from app.db.database import connect
from app.db.embedding_repository import EmbeddingRepository
from app.main import create_app
from app.providers.embedding_client import EmbeddingCallError, EmbeddingDimensionError

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx(paragraphs: list[str]) -> bytes:
    """造一份带标题的 DOCX（与知识库入库测试同一套夹具口径）。"""
    document = Document()
    document.add_heading("知识库资料", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, name: str, payload: bytes) -> dict:
    response = client.post("/api/knowledge/sources", files={"file": (name, payload, DOCX_MIME)})
    assert response.status_code == 201, response.text
    return response.json()


def _vectors(db_path, document_id: int) -> list[dict]:
    """直接查库读向量行（清单接口不暴露向量，这是第二处证据）。"""
    with connect(db_path) as connection:
        rows = connection.execute(
            """
            SELECT e.chunk_id, e.model_name, e.version, e.vector_json, c.document_id
              FROM embeddings e
              JOIN document_chunks c ON c.id = e.chunk_id
             WHERE c.document_id = ?
             ORDER BY c.seq ASC
            """,
            (document_id,),
        ).fetchall()
    return [dict(row) for row in rows]


def test_upload_writes_one_vector_per_chunk(client: TestClient, db_path) -> None:
    """上传成功后每个分块恰好一条向量，写入格式固定（JSON 浮点数组）。"""
    item = _upload(
        client,
        "带索引的资料.docx",
        _docx(["第一段用于验证向量写入。", "第二段同样需要一条向量。"]),
    )["item"]

    assert item["chunk_count"] > 0
    rows = _vectors(db_path, item["document_id"])
    assert len(rows) == item["chunk_count"]
    assert [row["chunk_id"] for row in rows]
    assert all(row["model_name"] == "fake-embedding" for row in rows)
    assert all(row["version"] == "v1" for row in rows)

    for row in rows:
        parsed = json.loads(row["vector_json"])
        # 固定格式：紧凑 JSON 数组（不带空格），元素全是浮点。
        assert row["vector_json"] == json.dumps(parsed, separators=(",", ":"))
        assert isinstance(parsed, list)
        assert parsed
        assert all(isinstance(value, float) for value in parsed)


def test_same_text_gets_same_vector(client: TestClient, db_path) -> None:
    """同一段文本无论上传几次都得到同一向量（确定性假向量，测试可复现）。"""
    text = "这一段文字会在两份资料里重复出现，用来验证向量确定性。"
    first = _upload(client, "甲.docx", _docx([text]))["item"]
    second = _upload(client, "乙.docx", _docx([text]))["item"]

    first_rows = _vectors(db_path, first["document_id"])
    second_rows = _vectors(db_path, second["document_id"])
    assert first_rows and second_rows
    assert json.loads(first_rows[0]["vector_json"]) == json.loads(second_rows[0]["vector_json"])


def test_embedding_failure_keeps_source_readable(
    client: TestClient, db_path, embedding_client: FakeEmbeddingClient
) -> None:
    """嵌入失败不丢资料：仍 201、分块仍在、状态「索引失败」且原因可读。"""
    embedding_client.set_error(EmbeddingCallError("嵌入模型未配置"))

    payload = _upload(client, "嵌入会失败.docx", _docx(["这段文字照样要能被读到。"]))
    item = payload["item"]

    assert item["index_status"] == "failed"
    assert item["index_status_label"] == "索引失败"
    assert item["index_message"]
    assert "向量索引失败" in item["index_message"]
    assert payload["message"] == item["index_message"]
    # 资料本身仍可读：分块落库了，只是没有向量。
    assert item["chunk_count"] > 0
    assert _vectors(db_path, item["document_id"]) == []

    listed = client.get("/api/knowledge/sources").json()["items"]
    assert [entry["source_id"] for entry in listed] == [item["source_id"]]
    assert listed[0]["chunk_count"] == item["chunk_count"]
    assert listed[0]["index_status"] == "failed"


def test_ragged_vectors_fail_with_dimension_message(
    client: TestClient, db_path, embedding_client: FakeEmbeddingClient
) -> None:
    """向量维度不一致时报错，且一条向量都不落库（不留半份索引）。"""
    embedding_client.set_ragged()

    item = _upload(
        client,
        "维度不一致.docx",
        _docx(["第一段内容。", "第二段内容需要多切几块以触发多条向量。" * 200]),
    )["item"]

    assert item["chunk_count"] > 1
    assert item["index_status"] == "failed"
    assert "维度" in (item["index_message"] or "")
    assert _vectors(db_path, item["document_id"]) == []


def test_reindex_after_model_switch_replaces_vectors(
    client: TestClient, db_path, embedding_client: FakeEmbeddingClient
) -> None:
    """换嵌入模型后重建索引：旧向量被整份替换，不与新向量混存。"""
    item = _upload(client, "会换模型.docx", _docx(["重建索引前的正文内容。"]))["item"]
    before = _vectors(db_path, item["document_id"])
    assert before and before[0]["model_name"] == "fake-embedding"

    embedding_client.set_model("another-embedding", dimensions=6)
    body = client.post("/api/knowledge/reindex", json={"source_id": item["source_id"]}).json()

    assert body["items"][0]["index_status"] == "stored"
    after = _vectors(db_path, item["document_id"])
    assert len(after) == item["chunk_count"]
    assert {row["model_name"] for row in after} == {"another-embedding"}
    assert {len(json.loads(row["vector_json"])) for row in after} == {6}
    # 旧模型的行一条都不剩（否则就是新旧混存）。
    assert not any(row["model_name"] == "fake-embedding" for row in after)


def test_reindex_marks_failure_when_embedding_breaks(
    client: TestClient, db_path, embedding_client: FakeEmbeddingClient
) -> None:
    """重建时嵌入失败：逐条如实回报失败，资料与分块仍在，向量为空。"""
    item = _upload(client, "重建会失败.docx", _docx(["重建索引时嵌入会挂掉。"]))["item"]

    embedding_client.set_error(EmbeddingCallError("嵌入服务不可用"))
    body = client.post("/api/knowledge/reindex", json={"source_id": item["source_id"]}).json()

    assert body["items"][0]["index_status"] == "failed"
    assert "向量索引失败" in (body["items"][0]["message"] or "")
    assert _vectors(db_path, item["document_id"]) == []

    listed = client.get("/api/knowledge/sources").json()["items"][0]
    assert listed["chunk_count"] > 0
    assert listed["index_status"] == "failed"


def test_reader_rejects_mixed_dimensions(client: TestClient, db_path) -> None:
    """读取端校验维度：有人手工改坏一行，读取必须报错而不是算出相似度。"""
    item = _upload(
        client,
        "会被改坏.docx",
        _docx(["这一段的向量稍后会被篡改。" * 200, "再来一段，确保切成多块。" * 200]),
    )["item"]
    rows = _vectors(db_path, item["document_id"])
    assert len(rows) > 1

    with connect(db_path) as connection:
        connection.execute(
            "UPDATE embeddings SET vector_json = ? WHERE id = (SELECT MIN(id) FROM embeddings)",
            (json.dumps([0.1, 0.2, 0.3]),),
        )
        connection.commit()

    with pytest.raises(EmbeddingDimensionError) as excinfo:
        EmbeddingRepository(db_path).list_vectors(item["document_id"])
    assert "维度" in str(excinfo.value)


def test_reader_rejects_mixed_models(client: TestClient, db_path) -> None:
    """读取端校验模型：换模型后未重建就把旧向量算进去，必须报错。"""
    item = _upload(
        client,
        "混了模型.docx",
        _docx(["第一段。", "第二段内容。" * 200]),
    )["item"]
    rows = _vectors(db_path, item["document_id"])
    assert len(rows) > 1

    with connect(db_path) as connection:
        connection.execute(
            "UPDATE embeddings SET model_name = 'legacy-embedding' WHERE id = (SELECT MIN(id) FROM embeddings)"
        )
        connection.commit()

    with pytest.raises(EmbeddingDimensionError) as excinfo:
        EmbeddingRepository(db_path).list_vectors(item["document_id"])
    assert "模型" in str(excinfo.value)


def test_embedding_seam_defaults_to_configured_client_and_can_be_overridden(
    db_path, data_dir, templates_dir
) -> None:
    """嵌入接缝：默认真实实现按环境变量构造，仍可被替身覆盖。

    这里刻意不用 `app` 夹具：那个夹具已把嵌入客户端换成替身，而本测试要看的是
    `create_app` 出厂状态下的默认接线。
    """
    application = create_app(
        db_path=db_path,
        mount_frontend=False,
        data_dir=data_dir,
        templates_dir=templates_dir,
    )

    @application.get("/api/_embedding_probe")
    def probe(embedding_client: EmbeddingClientDep) -> dict[str, str]:
        return {
            "embedding_client": type(embedding_client).__name__,
            "model_name": getattr(embedding_client, "model_name", ""),
        }

    with TestClient(application) as client:
        body = client.get("/api/_embedding_probe").json()
        assert body["embedding_client"] == "ConfiguredEmbeddingClient"
        assert body["model_name"] == "text-embedding-v4"

    fake = FakeEmbeddingClient(model_name="probe-embedding")
    application.dependency_overrides[get_embedding_client] = lambda: fake
    with TestClient(application) as client:
        body = client.get("/api/_embedding_probe").json()
    assert body["embedding_client"] == "FakeEmbeddingClient"
    assert body["model_name"] == "probe-embedding"
