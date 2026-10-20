"""总结用模板生成 Word（US16）接缝测试：默认 / 自定义模板、记录、下载与负测。"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from app.db.summary_deliverable_repository import SummaryDeliverableRepository

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

SUMMARY_FIELDS = {
    "id",
    "source_id",
    "task_id",
    "format",
    "file_path",
    "template_used",
    "created_at",
    "template_style_applied",
    "missing_styles",
    "headings",
}


def _short_docx() -> bytes:
    document = Document()
    document.add_heading("短资料", level=1)
    document.add_paragraph("这是一份用于总结 Word 成品测试的短资料。")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient) -> dict:
    response = client.post(
        "/api/analysis/sources",
        files={"file": ("短资料.docx", _short_docx(), DOCX_MIME)},
    )
    assert response.status_code == 201
    return response.json()


def _analyzed_source(client: TestClient) -> dict:
    uploaded = _upload(client)
    analyzed = client.post(f"/api/analysis/sources/{uploaded['source']['id']}/analyze")
    assert analyzed.status_code == 200
    return uploaded


def _custom_template(templates_dir: Path, name: str = "corp.docx") -> None:
    templates_dir.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_paragraph("公司模板封面")
    document.save(str(templates_dir / name))


def test_default_template_generates_word_deliverable_and_records_it(
    analysis_client: TestClient,
) -> None:
    """无 template 时走全局默认（无登记则内置模板），产物落盘并记录。"""
    uploaded = _analyzed_source(analysis_client)

    response = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/summary/word", json={}
    )

    assert response.status_code == 200
    body = response.json()
    assert set(body) == SUMMARY_FIELDS
    assert body["source_id"] == uploaded["source"]["id"]
    assert body["task_id"] == uploaded["source"]["task_id"]
    assert body["format"] == "docx"
    assert body["template_used"] == "_builtin/minimal-default.docx"
    assert body["template_style_applied"] is True
    assert body["headings"] >= 4

    artifact = Path(body["file_path"])
    assert artifact.exists()
    assert artifact.suffix == ".docx"
    assert artifact.parent.name == str(uploaded["source"]["task_id"])

    document = Document(str(artifact))
    texts = [paragraph.text for paragraph in document.paragraphs]
    assert "资料总结" in texts
    assert "来源" in texts
    assert "总结" in texts
    assert "引用" in texts


def test_custom_template_is_used_and_recorded(
    analysis_client: TestClient, templates_dir: Path
) -> None:
    """请求指定模板优先，并记录实际所用模板。"""
    _custom_template(templates_dir)
    uploaded = _analyzed_source(analysis_client)

    response = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/summary/word",
        json={"template": "corp.docx"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["template_used"] == "corp.docx"
    assert Path(body["file_path"]).exists()


def test_missing_summary_returns_409(analysis_client: TestClient) -> None:
    """尚无总结时拒绝生成（409）。"""
    uploaded = _upload(analysis_client)

    response = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/summary/word", json={}
    )

    assert response.status_code == 409
    assert "尚无总结" in response.json()["detail"]


def test_missing_template_returns_400(analysis_client: TestClient) -> None:
    """模板缺失明确报错（400），不静默退回内置模板。"""
    uploaded = _analyzed_source(analysis_client)

    response = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/summary/word",
        json={"template": "missing.docx"},
    )

    assert response.status_code == 400
    assert "模板" in response.json()["detail"]


def test_unknown_source_returns_404(analysis_client: TestClient) -> None:
    """来源不存在 404。"""
    assert analysis_client.post("/api/analysis/sources/9999/summary/word", json={}).status_code == 404


def test_generated_deliverable_can_be_downloaded(analysis_client: TestClient) -> None:
    """按成品记录 id 下载 .docx；不存在 404。"""
    uploaded = _analyzed_source(analysis_client)
    created = analysis_client.post(
        f"/api/analysis/sources/{uploaded['source']['id']}/summary/word", json={}
    ).json()

    download = analysis_client.get(f"/api/analysis/summary-deliverables/{created['id']}")

    assert download.status_code == 200
    assert download.content[:2] == b"PK"
    assert "wordprocessingml" in download.headers["content-type"]

    missing = analysis_client.get("/api/analysis/summary-deliverables/9999")
    assert missing.status_code == 404


def test_download_rejects_path_outside_deliverables(
    analysis_client: TestClient, db_path: Path, tmp_path: Path
) -> None:
    """下载端点只允许 data/deliverables/ 下的产物，拒绝库内越界路径。"""
    uploaded = _upload(analysis_client)
    outside = tmp_path / "outside.docx"
    outside.write_bytes(b"PK\x03\x04")
    repository = SummaryDeliverableRepository(db_path)
    record = repository.create(
        uploaded["source"]["id"],
        task_id=uploaded["source"]["task_id"],
        file_path=str(outside),
        template_used="outside.docx",
    )

    response = analysis_client.get(
        f"/api/analysis/summary-deliverables/{record['id']}"
    )

    assert response.status_code == 404
    assert "路径" in response.json()["detail"]