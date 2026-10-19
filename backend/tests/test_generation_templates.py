"""TB-5 模板机制与负测（唯一接缝 = 进程内 HTTP；模型走 fake；不联网）。

覆盖工单 14 的关键决策与两个已核实的硬陷阱（调研 A2）：
- 内置最小默认模板保证开箱可用；用户模板放进 `templates/` 即可选用；
- **全局默认模板**与**任务级临时覆盖**在数据与 API 上互不混淆，每次生成记录实际所用模板；
- `.dotx` / `.dotm` / `.docm`（含改名为 `.docx` 的模板类型文件）被拒并提示「请另存为 .docx」；
- 模板样式未定义 → 两段式回退 + `template_style_applied: false` 标记；
- 模板缺失 / 损坏 → 明确报错，**不产生成品文件**，绝不静默退回空白生成；
- 模板名不得越出 `templates/` 白名单。
"""

from __future__ import annotations

import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.shared import Pt
from fakes import ScriptedModelClient
from fastapi.testclient import TestClient

from app.generation.service import verify_word_deliverable
from app.generation.templates import ensure_builtin_template
from app.generation.word import DeliverableIntegrityError, build_word_deliverable

IDEA = "帮我准备一份面向新员工的产品培训"

# 短方案：H1 / H2 / H3 + 正文 + 列表，够验证样式映射与回退分支。
PLAN_MD = (
    "# 方案\n\n"
    "## 目标\n掌握产品全貌。\n\n"
    "### 定位\n从场景切入。\n\n"
    "- 早晨开工\n"
)

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)


def _frontier_json(*triples: tuple[str, str, str]) -> str:
    frontier = [
        {"title": title, "body": body, "recommended_answer": answer}
        for title, body, answer in triples
    ]
    return json.dumps({"frontier": frontier}, ensure_ascii=False)


def _confirmed_task(client: TestClient, scripted: ScriptedModelClient) -> int:
    """走完拷问 → 收敛 → 确认方案，返回已确认方案的任务 id。"""
    scripted.set_sequence(
        _frontier_json(("Q", "正文", "推荐")),
        json.dumps({"frontier": []}, ensure_ascii=False),
        PLAN_MD,
    )
    task_id = client.post("/api/grilling/tasks", json={"idea": IDEA}).json()["task"]["id"]
    client.post(
        f"/api/grilling/tasks/{task_id}/answers", json={"round_no": 1, "answers": ["答"]}
    )
    client.post(f"/api/grilling/tasks/{task_id}/advance")
    assert client.post(f"/api/grilling/tasks/{task_id}/plan/confirm").status_code == 200
    return task_id


def _make_user_template(path: Path, *, heading1_pt: int | None = None) -> Path:
    """造一份「用户模板」（可调 Heading 1 字号，用于验证样式随模板变化）。"""
    document = Document()
    if heading1_pt is not None:
        document.styles["Heading 1"].font.size = Pt(heading1_pt)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def _without_styles(src: Path, dst: Path, ui_names: tuple[str, ...]) -> Path:
    """复制一份 .docx 并删掉指定样式的实体定义（模拟调研 A2 陷阱 2）。"""
    from docx.styles import BabelFish

    internal = {BabelFish.ui2internal(name) for name in ui_names}
    with zipfile.ZipFile(src) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    root = ElementTree.fromstring(parts["word/styles.xml"])
    w = f"{{{_W_NS}}}"
    for style in list(root.findall(f"{w}style")):
        name_element = style.find(f"{w}name")
        if name_element is not None and name_element.get(f"{w}val") in internal:
            root.remove(style)
    parts["word/styles.xml"] = ElementTree.tostring(
        root, xml_declaration=True, encoding="UTF-8"
    )

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return dst


def _as_template_typed(src: Path, dst: Path) -> Path:
    """把 .docx 的主部件内容类型改成「模板」——即改名为 .docx 的 .dotx（陷阱 1）。"""
    with zipfile.ZipFile(src) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    root = ElementTree.fromstring(parts["[Content_Types].xml"])
    for override in root.findall(f"{{{_CT_NS}}}Override"):
        if override.get("PartName") == "/word/document.xml":
            override.set("ContentType", _TEMPLATE_CONTENT_TYPE)
    parts["[Content_Types].xml"] = ElementTree.tostring(
        root, xml_declaration=True, encoding="UTF-8"
    )

    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return dst


def _deliverable_files(data_dir: Path) -> list[Path]:
    root = data_dir / "deliverables"
    return sorted(root.rglob("*.docx")) if root.is_dir() else []


# --- 内置最小默认模板：开箱可用 -------------------------------------------------


def test_builtin_default_template_is_usable_without_registration(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """未放任何用户模板时也能生成：内置最小默认模板兜底，并在清单里可见。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)

    generated = grilling_client.post(f"/api/tasks/{task_id}/generations")

    assert generated.status_code == 201
    assert generated.json()["template_used"] == "_builtin/minimal-default.docx"

    listed = grilling_client.get("/api/templates").json()["items"]
    builtin = next(item for item in listed if item["name"] == "_builtin/minimal-default.docx")
    assert builtin["builtin"] is True
    assert builtin["usable"] is True
    assert builtin["format"] == "docx"


# --- 负测：模板缺失 / 损坏，绝不静默退回空白生成 --------------------------------


def test_missing_template_reports_error_and_produces_no_deliverable(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient, data_dir: Path
) -> None:
    """指定了不存在的模板 → 明确报错；不产生成品文件，也不留成品记录。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)

    response = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "not-there.docx"}
    )

    assert response.status_code == 400
    assert "不存在" in response.json()["detail"]
    assert grilling_client.get(f"/api/tasks/{task_id}/generations").json()["total"] == 0
    assert _deliverable_files(data_dir) == []  # 绝不静默退回空白生成


def test_missing_global_default_template_reports_error(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
    data_dir: Path,
) -> None:
    """全局默认模板登记后文件被移走 → 报错，而不是悄悄改用内置模板。"""
    _make_user_template(templates_dir / "corp.docx")
    assert grilling_client.post("/api/templates", json={"name": "corp.docx"}).status_code == 201
    assert grilling_client.put("/api/templates/corp.docx/default").status_code == 200
    (templates_dir / "corp.docx").unlink()

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    response = grilling_client.post(f"/api/tasks/{task_id}/generations")

    assert response.status_code == 400
    assert "全局默认模板" in response.json()["detail"]
    assert _deliverable_files(data_dir) == []


def test_dotx_template_is_rejected_with_save_as_docx_hint(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
    data_dir: Path,
) -> None:
    """`.dotx` 一律不被接受：登记与临时覆盖都给「请另存为 .docx」的友好错误。"""
    _make_user_template(templates_dir / "corp.dotx")

    registered = grilling_client.post("/api/templates", json={"name": "corp.dotx"})
    assert registered.status_code == 400
    assert "另存为 .docx" in registered.json()["detail"]

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    generated = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "corp.dotx"}
    )
    assert generated.status_code == 400
    assert "另存为 .docx" in generated.json()["detail"]
    assert _deliverable_files(data_dir) == []


def test_template_typed_file_renamed_to_docx_is_rejected(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
) -> None:
    """陷阱 1 的真身：把模板类型文件改名成 `.docx`，python-docx 打不开，同样被拒。"""
    source = _make_user_template(templates_dir / "corp-source.docx")
    _as_template_typed(source, templates_dir / "corp.docx")

    registered = grilling_client.post("/api/templates", json={"name": "corp.docx"})

    assert registered.status_code == 400
    detail = registered.json()["detail"]
    assert "另存为 .docx" in detail
    assert "Word 文档打开" in detail

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    generated = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "corp.docx"}
    )
    assert generated.status_code == 400


def test_template_name_cannot_escape_templates_dir(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    data_dir: Path,
    templates_dir: Path,
) -> None:
    """模板名不得越出 `templates/` 白名单（路径遍历被拒）。"""
    (templates_dir.parent / "secret.docx").write_bytes(b"secret")

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    response = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "../secret.docx"}
    )

    assert response.status_code == 400
    assert "越出" in response.json()["detail"]
    assert _deliverable_files(data_dir) == []


# --- 陷阱 2：模板样式未定义 → 两段式回退且降级可见 ------------------------------


def test_template_without_required_styles_falls_back_and_flags_it(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
) -> None:
    """模板缺 Heading/List 样式 → 回退为直接格式，并标 `template_style_applied: false`。"""
    source = _make_user_template(templates_dir / "plain-source.docx")
    _without_styles(
        source,
        templates_dir / "plain.docx",
        ("Heading 1", "Heading 2", "Heading 3", "List Bullet"),
    )

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    response = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "plain.docx"}
    )

    assert response.status_code == 201
    body = response.json()
    assert body["template_style_applied"] is False
    assert set(body["missing_styles"]) >= {"Heading 1", "Heading 2", "Heading 3"}
    assert body["integrity_ok"] is True
    assert body["template_used"] == "plain.docx"

    document = Document(body["file_path"])
    # 没有 Heading 样式可用：标题必须是「加粗的直接格式段落」
    assert {p.style.name for p in document.paragraphs} == {"Normal"}
    bold_texts = {p.text for p in document.paragraphs if p.runs and p.runs[0].bold}
    assert {"方案", "目标", "定位"} <= bold_texts
    # 列表回退为带符号前缀的普通段落，内容不丢
    assert any(p.text.startswith("• 早晨开工") for p in document.paragraphs)


def test_template_with_all_styles_keeps_template_appearance(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
) -> None:
    """样式齐备时按模板样式写：`template_style_applied` 为 true，标题走 `Heading N`。"""
    _make_user_template(templates_dir / "corp.docx", heading1_pt=28)

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    body = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "corp.docx"}
    ).json()

    assert body["template_style_applied"] is True
    assert body["missing_styles"] == []

    document = Document(body["file_path"])
    assert document.styles["Heading 1"].font.size == Pt(28)  # 模板外观被保留
    assert document.paragraphs[0].style.name == "Heading 1"


# --- 全局默认模板 vs 任务级临时覆盖：互不混淆 -----------------------------------


def test_global_default_and_task_override_stay_separate(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    templates_dir: Path,
) -> None:
    """临时覆盖只影响单次生成，不改变全局默认；两者每次都被如实记录。"""
    _make_user_template(templates_dir / "red.docx", heading1_pt=28)
    _make_user_template(templates_dir / "blue.docx", heading1_pt=14)
    assert grilling_client.post("/api/templates", json={"name": "red.docx"}).status_code == 201
    assert grilling_client.put("/api/templates/red.docx/default").status_code == 200

    task_id = _confirmed_task(grilling_client, scripted_model_client)

    by_default = grilling_client.post(f"/api/tasks/{task_id}/generations").json()
    assert by_default["template_used"] == "red.docx"

    override = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"template": "blue.docx"}
    ).json()
    assert override["template_used"] == "blue.docx"
    assert Document(override["file_path"]).styles["Heading 1"].font.size == Pt(14)
    assert Document(by_default["file_path"]).styles["Heading 1"].font.size == Pt(28)

    # 临时覆盖没有改写全局默认：清单里默认仍是 red，且下一次生成仍走 red
    listed = {item["name"]: item for item in grilling_client.get("/api/templates").json()["items"]}
    assert listed["red.docx"]["is_default"] is True
    assert listed["blue.docx"]["is_default"] is False

    again = grilling_client.post(f"/api/tasks/{task_id}/generations").json()
    assert again["template_used"] == "red.docx"

    # 每次生成都留下自己的记录（含实际所用模板），可复现
    records = grilling_client.get(f"/api/tasks/{task_id}/generations").json()["items"]
    assert [record["template_used"] for record in records] == [
        "red.docx",
        "blue.docx",
        "red.docx",
    ]


def test_registration_validates_file_and_sets_global_default(
    grilling_client: TestClient, templates_dir: Path
) -> None:
    """登记：不存在的文件 404；已放入的 .docx 可登记、可设为全局默认。"""
    assert grilling_client.post("/api/templates", json={"name": "ghost.docx"}).status_code == 404

    _make_user_template(templates_dir / "corp.docx")
    registered = grilling_client.post("/api/templates", json={"name": "corp.docx"})

    assert registered.status_code == 201
    assert registered.json()["template"]["format"] == "docx"
    assert registered.json()["template"]["registered"] is True
    assert registered.json()["default_template"] is None  # 未指定默认时不自动成为默认

    made_default = grilling_client.put("/api/templates/corp.docx/default")

    assert made_default.status_code == 200
    assert made_default.json()["template"]["is_default"] is True
    assert made_default.json()["default_template"] == "corp.docx"


def test_registered_template_that_vanished_is_reported_as_missing(
    grilling_client: TestClient, templates_dir: Path
) -> None:
    """已登记但文件不在：如实标 `missing`，不假装模板还在。"""
    path = _make_user_template(templates_dir / "corp.docx")
    grilling_client.post("/api/templates", json={"name": "corp.docx"})
    path.unlink()

    listed = grilling_client.get("/api/templates").json()["items"]

    entry = next(item for item in listed if item["name"] == "corp.docx")
    assert entry["missing"] is True
    assert entry["usable"] is False


# --- 完整性检查（规格 US 22）---------------------------------------------------


def test_integrity_check_rejects_corrupted_artifact(tmp_path: Path) -> None:
    """产物结构损坏（截断的 zip）→ 完整性检查报错，不交付。"""
    template = ensure_builtin_template(tmp_path / "templates")
    output = tmp_path / "out" / "plan.docx"

    build = build_word_deliverable(PLAN_MD, template, output)
    verify_word_deliverable(build)  # 合格产物先通过

    raw = output.read_bytes()
    output.write_bytes(raw[: len(raw) // 2])

    with pytest.raises(DeliverableIntegrityError):
        verify_word_deliverable(build)


def test_broken_artifact_is_discarded_and_not_recorded(
    grilling_client: TestClient,
    scripted_model_client: ScriptedModelClient,
    data_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """完整性检查未通过时不交付：产物被丢弃、不落库，客户端拿到明确的 5xx。"""

    def _always_broken(build) -> None:
        raise DeliverableIntegrityError("模拟：章节数不符")

    monkeypatch.setattr("app.generation.service.verify_word_deliverable", _always_broken)

    task_id = _confirmed_task(grilling_client, scripted_model_client)
    response = grilling_client.post(f"/api/tasks/{task_id}/generations")

    assert response.status_code == 500
    assert "完整性检查未通过" in response.json()["detail"]
    assert grilling_client.get(f"/api/tasks/{task_id}/generations").json()["total"] == 0
    assert _deliverable_files(data_dir) == []


# --- 另一轨：单文件 HTML 演示由 TB-6 落地 ---------------------------------------


def test_html_format_is_not_implemented_yet(
    grilling_client: TestClient, scripted_model_client: ScriptedModelClient
) -> None:
    """ADR-0003 的另一轨（单文件 HTML 演示）显式 501，不假装已实现。"""
    task_id = _confirmed_task(grilling_client, scripted_model_client)

    response = grilling_client.post(
        f"/api/tasks/{task_id}/generations", json={"format": "html"}
    )

    assert response.status_code == 501